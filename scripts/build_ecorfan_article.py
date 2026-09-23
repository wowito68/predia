#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
TMP = Path("/tmp/predia_article_work")
TEMPLATE = Path("/home/wowo/Descargas/c-Journal-Template-I-4.docx")
ARTICLE_TEXT = TMP / "article.txt"
BBL = ROOT / "paper" / "paper_ml_diabetes_mimeti2026_es.bbl"
FINAL_DOCX = ROOT / "docs" / "predia-articulo-ecorfan-ciermmi-2026.docx"
GRAPHICAL = TMP / "graphical_abstract.png"

TNR = "Times New Roman"
ACCENT = RGBColor(32, 55, 100)
LIGHT_FILL = "EAF1F8"
HEADER_FILL = "1F4E79"
GRAY_FILL = "F2F2F2"


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def clear_paragraphs(container) -> None:
    for p in list(container.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(container.tables):
        t._element.getparent().remove(t._element)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="BFBFBF", size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color="7F7F7F", size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)


def set_columns(section, num: int) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "360")


def fmt_run(run, size=10.5, bold=False, italic=False, color=None, all_caps=False):
    run.font.name = TNR
    run._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if all_caps:
        run.font.all_caps = True
    return run


def add_p(doc, text="", size=11, bold=False, italic=False, align=None, color=None, before=0, after=3, line=1.0):
    p = doc.add_paragraph()
    if text:
        fmt_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    size = 12 if level == 1 else 10.5
    p = add_p(doc, text, size=size, bold=True, color=ACCENT if level == 1 else None, before=8 if level == 1 else 5, after=3)
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, label, text):
    p = add_p(doc, "", size=8.5, after=4)
    fmt_run(p.add_run(label), size=8.5, bold=True)
    fmt_run(p.add_run(f" {text}"), size=8.5)
    return p


def add_note_box(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    apply_table_widths(table, [3.0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    set_cell_border(cell, color="A6BDD7", size="4")
    set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    fmt_run(p.add_run(title), size=9.5, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    fmt_run(p2.add_run(text), size=9)
    add_p(doc, "", size=1, after=2)


def apply_table_widths(table, widths):
    if not widths:
        defaults = {
            2: [1.25, 1.65],
            3: [0.75, 1.35, 0.85],
            5: [0.9, 0.55, 0.55, 0.5, 0.5],
        }
        widths = defaults.get(len(table.columns), [3.0 / len(table.columns)] * len(table.columns))
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(Inches(w).twips for w in widths)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(Inches(width).twips))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = Inches(widths[idx])
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width.twips))


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    apply_table_widths(table, widths)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        set_cell_shading(hdr[i], HEADER_FILL)
        set_cell_border(hdr[i], color="4F81BD", size="4")
        set_cell_margins(hdr[i], top=70, start=70, bottom=70, end=70)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_run(p.add_run(str(h)), size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            set_cell_border(cells[i], color="D9D9D9", size="3")
            set_cell_margins(cells[i], top=70, start=70, bottom=70, end=70)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            fmt_run(p.add_run(str(value)), size=font_size)
    apply_table_widths(table, widths)
    add_p(doc, "", size=1, after=4)
    return table


def create_graphical_abstract() -> None:
    w, h = 1480, 620
    img = Image.new("RGB", (w, h), "#ffffff")
    d = ImageDraw.Draw(img)
    try:
        font_bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 38)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 29)
        small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
    except Exception:
        font_bold = font = small = None
    d.rectangle([0, 0, w - 1, h - 1], outline="#1f4e79", width=3)
    d.rectangle([0, 0, w, 82], fill="#1f4e79")
    d.text((42, 24), "Benchmark interpretable para riesgo de diabetes", fill="#ffffff", font=font_bold)
    boxes = [
        (55, 135, 425, 500, "#EAF1F8", "Datos", ["BRFSS 2015", "253,680 personas", "21 indicadores poblacionales"]),
        (555, 135, 925, 500, "#E2F0D9", "Modelado", ["9 estimadores", "validacion 60/20/20", "calibracion y SHAP"]),
        (1055, 135, 1425, 500, "#FFF2CC", "Hallazgo", ["AUC max. 0.827", "+0.007 vs regresion", "cribado, no diagnostico"]),
    ]
    for x1, y1, x2, y2, fill, title, lines in boxes:
        d.rounded_rectangle([x1, y1, x2, y2], radius=26, fill=fill, outline="#A6A6A6", width=3)
        d.text((x1 + 34, y1 + 34), title, fill="#203864", font=font_bold)
        y = y1 + 115
        for line in lines:
            d.ellipse([x1 + 36, y + 7, x1 + 54, y + 25], fill="#1f4e79")
            d.text((x1 + 72, y), line, fill="#222222", font=font)
            y += 58
    d.line([430, 318, 545, 318], fill="#1f4e79", width=8)
    d.polygon([(545, 318), (515, 300), (515, 336)], fill="#1f4e79")
    d.line([930, 318, 1045, 318], fill="#1f4e79", width=8)
    d.polygon([(1045, 318), (1015, 300), (1015, 336)], fill="#1f4e79")
    d.text((55, 548), "Uso propuesto: priorizacion clinica y educacion del paciente, con validacion medica posterior.", fill="#595959", font=small)
    img.save(GRAPHICAL)


def clean_latex_reference(text: str) -> str:
    text = re.sub(r"%.*", "", text)
    text = text.replace("~", " ")
    text = text.replace("---", " - ")
    text = text.replace("``", '"').replace("''", '"')
    text = text.replace("\\newblock", " ")
    text = re.sub(r"\\BIBentry[A-Za-z]+", " ", text)
    text = re.sub(r"\\hskip.*?\\relax", " ", text)
    text = re.sub(r"\\url\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\emph\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\BIBforeignlanguage\{[^{}]+\}\{([^{}]+)\}", r"\1", text)
    text = text.replace("{\\'E}", "É").replace("{\\v{s}}", "š")
    text = re.sub(r"\\[a-zA-Z]+(?:\[[^\]]+\])?", " ", text)
    text = text.replace("{", "").replace("}", "").replace("\\", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,.;")


def parse_bbl_references():
    if not BBL.exists():
        return []
    text = BBL.read_text(encoding="utf-8", errors="ignore")
    if "\\begin{thebibliography}" not in text:
        return []
    chunks = text.split("\\bibitem")[1:]
    refs = []
    for idx, chunk in enumerate(chunks, start=1):
        chunk = re.sub(r"^\{[^{}]+\}", "", chunk, count=1)
        chunk = chunk.split("\\end{thebibliography}")[0]
        cleaned = clean_latex_reference(chunk)
        if cleaned:
            refs.append(f"[{idx}] {cleaned}.")
    return refs


def parse_pdf_references():
    text = ARTICLE_TEXT.read_text(encoding="utf-8", errors="ignore")
    marker = "R EFERENCIAS"
    start = text.find(marker)
    if start == -1:
        return []
    lines = text[start + len(marker) :].splitlines()
    refs, cur = [], []
    for line in lines:
        s = " ".join(line.split())
        if re.match(r"^\[\d+\]", s):
            if cur:
                refs.append(" ".join(cur))
            cur = [s]
        elif cur and s:
            cur.append(s)
    if cur:
        refs.append(" ".join(cur))
    return [re.sub(r"\s+", " ", ref).strip() for ref in refs if len(ref) > 15]


def parse_references():
    refs = parse_bbl_references()
    if refs:
        return refs[:35]
    return parse_pdf_references()[:35]


def load_model_rows():
    path = ROOT / "ml-research" / "brfss" / "results" / "models" / "comparison_test.csv"
    if not path.exists():
        return [
            ["HistGradientBoosting", "0.8268", "0.4233", "0.2461", "0.0974"],
            ["XGBoost", "0.8268", "0.4232", "0.2459", "0.0974"],
            ["Logistic Regression", "0.8197", "0.3925", "0.3558", "0.1020"],
        ]
    rows = []
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("model", "").replace("_", " ").title()
            rows.append(
                [
                    name,
                    f"{float(row.get('roc_auc', 0)):.4f}",
                    f"{float(row.get('pr_auc', 0)):.4f}",
                    f"{float(row.get('mcc', 0)):.4f}",
                    f"{float(row.get('brier', 0)):.4f}",
                ]
            )
    return rows[:9]


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(13)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.28)
        set_columns(section, 1)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0


def configure_header_footer(doc: Document, page_count_label="1-14") -> None:
    for section in doc.sections:
        section.different_first_page_header_footer = True

        for header in (section.header, section.first_page_header, section.even_page_header):
            header.is_linked_to_previous = False
            clear_paragraphs(header)
            p = header.add_paragraph()
            fmt_run(p.add_run("ECORFAN Journal"), size=14, bold=True, color=ACCENT)
            fmt_run(p.add_run(f"\tOctober, 2026 Vol.1 No.1 {page_count_label}-[Using ECORFAN]"), size=8.5, color=RGBColor(89, 89, 89))
            p.paragraph_format.space_after = Pt(0)
            add_bottom_border(p, color="7F7F7F", size="6")

        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            clear_paragraphs(footer)
            fp = footer.add_paragraph()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt_run(fp.add_run("ISSN-On line: Use Only ECORFAN | ECORFAN all rights reserved"), size=8, color=RGBColor(89, 89, 89))


def build_document(page_count_label="1-14") -> None:
    FINAL_DOCX.parent.mkdir(parents=True, exist_ok=True)
    create_graphical_abstract()
    refs = parse_references()

    work = TMP / "working-ecorfan.docx"
    shutil.copyfile(TEMPLATE, work)
    doc = Document(work)
    clear_body(doc)
    configure_page(doc)
    configure_styles(doc)
    configure_header_footer(doc, page_count_label=page_count_label)

    # Front matter
    add_p(doc, "Article", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
    title = (
        "¿Cuánto mejora realmente el aprendizaje automático la predicción del riesgo de diabetes a partir de "
        "indicadores de salud poblacional? Un benchmark honesto, interpretable y reproducible sobre BRFSS 2015"
    )
    add_p(doc, title, size=14, bold=True, color=ACCENT, after=6)
    add_p(
        doc,
        "How much does machine learning really improve diabetes risk prediction from population health indicators? "
        "An honest, interpretable and reproducible benchmark on BRFSS 2015",
        size=11,
        italic=True,
        after=5,
    )
    add_p(
        doc,
        "ÁLVAREZ-SÁNCHEZ, G.; MOYA-MOYA, J. Javier; PEREDO-VALDERRAMA, I.; PACINDO-PIÑA, J.; FERNANDEZ-CONDE, D.",
        size=10.5,
        bold=True,
        after=3,
    )
    add_p(
        doc,
        "Universidad Politécnica de Querétaro, ROR https://ror.org/0580rm578. "
        "G. Álvarez Sánchez: ORCID 0009-0005-6254-143X; "
        "ResearchID rid168626 (researchid.co/rid168626; usuario wowoso10); CVU CONAHCYT 2044833; "
        "correo institucional 124050514@upq.edu.mx. Datos de coautores: pendientes de confirmación.",
        size=9,
        italic=True,
        after=2,
    )
    add_p(
        doc,
        "Received: use only ECORFAN | Accepted: use only ECORFAN | DOI: use only ECORFAN | Correspondence: 124050514@upq.edu.mx",
        size=8.5,
        italic=True,
        after=5,
    )
    p = add_p(doc, "", after=4)
    add_bottom_border(p, color="7F7F7F", size="8")

    front = doc.add_table(rows=1, cols=2)
    front.autofit = True
    for cell in front.rows[0].cells:
        set_cell_border(cell, color="FFFFFF", size="0")
        set_cell_margins(cell, top=70, start=90, bottom=70, end=90)

    left = front.cell(0, 0)
    right = front.cell(0, 1)
    for cell in (left, right):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cell.text = ""

    def cell_heading(cell, text):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        fmt_run(p.add_run(text), size=11, bold=True, color=ACCENT)

    def cell_text(cell, text, size=9.5):
        p = cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)
        fmt_run(p.add_run(text), size=size)

    cell_heading(left, "Resumen")
    cell_text(
        left,
        "Se presenta un benchmark reproducible para estimar riesgo de diabetes usando BRFSS 2015 "
        "(n=253,680; prevalencia 13.9%). Se compararon nueve estimadores con partición 60/20/20, "
        "validación, calibración y análisis interpretable. HistGradientBoosting y XGBoost alcanzaron "
        "ROC-AUC 0.827, apenas 0.007 por encima de regresión logística (0.820), lo que muestra que la "
        "calidad de las variables limita más que la complejidad del modelo. La calibración isotónica "
        "redujo ECE a 0.003 y las bandas de riesgo separaron prevalencias de 2.9% a 48.0%. El sistema "
        "apoya cribado y priorización clínica, no diagnóstico automático.",
    )
    cell_heading(left, "Abstract")
    cell_text(
        left,
        "This paper presents a reproducible benchmark for diabetes risk prediction using BRFSS 2015 "
        "(n=253,680; 13.9% prevalence). Nine estimators were compared under a 60/20/20 split with "
        "validation, calibration and interpretability analysis. HistGradientBoosting and XGBoost reached "
        "ROC-AUC 0.827, only 0.007 above logistic regression (0.820), indicating that variable quality "
        "limits performance more than model complexity. Isotonic calibration reduced ECE to 0.003 and "
        "risk bands separated observed prevalence from 2.9% to 48.0%. The model is suitable for screening "
        "and clinical prioritization, not autonomous diagnosis.",
    )
    cell_heading(left, "Palabras clave")
    cell_text(left, "Predicción de diabetes; aprendizaje automático; calibración clínica.", size=9)

    cell_heading(right, "Graphical Abstract")
    right.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p = right.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(GRAPHICAL), width=Inches(3.05))
    cell_heading(right, "Clasificación SECIHTI")
    cell_text(
        right,
        "Área: Ciencias de la Salud. Campo: Salud pública e informática médica. "
        "Disciplina: aprendizaje automático aplicado a salud. Subdisciplina: predicción clínica del riesgo de diabetes.",
        size=9,
    )
    cell_heading(right, "Keywords")
    cell_text(right, "Diabetes risk prediction; machine learning; clinical calibration.", size=9)

    add_p(
        doc,
        "Citation: Álvarez-Sánchez, G., Moya-Moya, J. J., Peredo-Valderrama, I., Pacindo-Piña, J. & Fernandez-Conde, D. "
        "(2026). ¿Cuánto mejora realmente el aprendizaje automático la predicción del riesgo de diabetes a partir de "
        f"indicadores de salud poblacional? ECORFAN Journal, Vol.1 No.1, {page_count_label}.",
        size=8,
        italic=True,
        before=4,
        after=0,
    )

    # Body in two columns.
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.page_width = Inches(8.5)
    body_section.page_height = Inches(13)
    body_section.left_margin = Inches(0.79)
    body_section.right_margin = Inches(0.79)
    body_section.top_margin = Inches(0.79)
    body_section.bottom_margin = Inches(0.79)
    body_section.header_distance = Inches(0.28)
    body_section.footer_distance = Inches(0.28)
    set_columns(body_section, 2)
    configure_header_footer(doc, page_count_label=page_count_label)

    add_heading(doc, "1. Introducción")
    intro = [
        "La diabetes mellitus tipo 2 es un problema prioritario de salud pública por su prevalencia, costos asistenciales y asociación con complicaciones cardiovasculares, renales, neurológicas y oftalmológicas. La detección temprana permite intervenir con cambios de estilo de vida, seguimiento preventivo y tratamiento oportuno antes de que aparezcan daños irreversibles.",
        "En este contexto, los modelos de aprendizaje automático se han popularizado como herramientas para estratificar riesgo. Sin embargo, una mejora aparente de métricas puede ocultar problemas de fuga de información, desbalance de clases, falta de calibración y poca utilidad clínica. Por ello, este trabajo se enfoca en una pregunta práctica: cuánto aporta realmente el aprendizaje automático frente a modelos lineales cuando solo se dispone de indicadores poblacionales autodeclarados.",
        "El objetivo del estudio no es producir un diagnóstico automático, sino establecer una línea base honesta, interpretable y reproducible para orientar sistemas como PREDIA, donde las predicciones deben integrarse a flujos clínicos verificables, explicables y supervisados por profesionales de salud.",
    ]
    for text in intro:
        add_p(doc, text)

    add_heading(doc, "2. Trabajo relacionado")
    for text in [
        "La literatura reciente reporta desempeño alto en predicción de diabetes con árboles de gradiente, ensambles y redes neuronales. No obstante, parte de estos resultados depende de variables de laboratorio o de datos clínicos muy cercanos al diagnóstico, lo que puede incrementar el AUC pero limitar la generalización cuando el sistema se usa para tamizaje poblacional.",
        "Otros trabajos muestran que los modelos lineales permanecen competitivos cuando los predictores son sociodemográficos y de hábitos de salud. Esto justifica comparar modelos complejos contra regresión logística y no contra baselines débiles. Además, en salud digital resulta tan relevante estimar probabilidades calibradas como ordenar pacientes por riesgo.",
    ]:
        add_p(doc, text)

    add_heading(doc, "3. Descripción del conjunto de datos")
    add_p(
        doc,
        "Se utilizó el conjunto Diabetes Health Indicators del Behavioral Risk Factor Surveillance System 2015 (CDC BRFSS). El conjunto contiene indicadores poblacionales autodeclarados sobre salud general, hábitos, edad, sexo, índice de masa corporal, hipertensión, colesterol, actividad física y consumo de frutas/verduras, entre otros. La variable objetivo binaria indica presencia o ausencia de diabetes.",
    )
    add_caption(doc, "Tabla 1.", "Características principales del conjunto de datos BRFSS 2015.")
    add_table(
        doc,
        ["Elemento", "Valor"],
        [
            ["Personas encuestadas", "253,680"],
            ["Predictores numéricos", "21"],
            ["Variable objetivo", "Diabetes_binary"],
            ["Prevalencia positiva", "13.9%"],
            ["Relación de desbalance", "6.18:1"],
            ["Valores faltantes", "0"],
            ["Filas duplicadas", "24,206 (9.5%), preservadas"],
            ["Variables de laboratorio", "No disponibles"],
        ],
        font_size=8.4,
    )
    add_note_box(
        doc,
        "Interpretación clínica",
        "La ausencia de glucosa, HbA1c u otros biomarcadores reduce el techo predictivo. Por ello, el modelo debe verse como apoyo de cribado y priorización, no como sustituto diagnóstico.",
    )
    add_caption(doc, "Tabla 2.", "Predictores utilizados y lectura clínica general.")
    add_table(
        doc,
        ["Variable", "Descripción", "Lectura"],
        [
            ["HighBP", "Presión arterial alta", "factor cardiometabólico"],
            ["HighChol", "Colesterol alto", "factor cardiometabólico"],
            ["CholCheck", "Revisión de colesterol", "seguimiento preventivo"],
            ["BMI", "Índice de masa corporal", "riesgo antropométrico"],
            ["Smoker", "Tabaquismo", "hábito de riesgo"],
            ["Stroke", "Antecedente de EVC", "comorbilidad"],
            ["HeartDiseaseorAttack", "Cardiopatía o infarto", "comorbilidad"],
            ["PhysActivity", "Actividad física", "protección conductual"],
            ["Fruits", "Consumo de frutas", "hábito alimentario"],
            ["Veggies", "Consumo de verduras", "hábito alimentario"],
            ["HvyAlcoholConsump", "Consumo intenso de alcohol", "hábito de riesgo"],
            ["AnyHealthcare", "Acceso a atención médica", "acceso al sistema"],
            ["NoDocbcCost", "Sin médico por costo", "barrera económica"],
            ["GenHlth", "Salud general percibida", "estado subjetivo"],
            ["MentHlth", "Días de salud mental afectada", "bienestar"],
            ["PhysHlth", "Días de salud física afectada", "funcionalidad"],
            ["DiffWalk", "Dificultad para caminar", "movilidad"],
            ["Sex", "Sexo", "demografía"],
            ["Age", "Grupo de edad", "demografía"],
            ["Education", "Escolaridad", "determinante social"],
            ["Income", "Ingreso", "determinante social"],
        ],
        font_size=6.8,
    )

    add_heading(doc, "4. Metodología")
    for text in [
        "El flujo experimental se diseñó para minimizar optimismo estadístico. La base se dividió en entrenamiento, validación y prueba con proporción 60/20/20. La selección de hiperparámetros y umbrales se realizó sin tocar el conjunto de prueba, que quedó reservado para la estimación final.",
        "Se evaluaron nueve estimadores: regresión logística, Random Forest, Extra Trees, KNN, SVM, MLP, LightGBM, XGBoost e HistGradientBoosting. Se calcularon ROC-AUC, PR-AUC, Brier score, MCC, exactitud balanceada, F1, sensibilidad y especificidad. La calibración se comparó en versión original, Platt e isotónica.",
        "La interpretación se realizó con importancia global y análisis SHAP para identificar factores de riesgo consistentes con conocimiento clínico. Las bandas de riesgo se definieron por probabilidad predicha para facilitar lectura médica y comunicación con pacientes.",
    ]:
        add_p(doc, text)

    add_heading(doc, "4.1 Preparación y control de calidad", level=2)
    for text in [
        "Las variables se conservaron en formato numérico para mantener trazabilidad respecto al diccionario original del BRFSS. No se eliminaron observaciones duplicadas porque la base pública conserva respuestas agregadas con combinaciones repetidas de factores; retirarlas sin una regla epidemiológica explícita habría alterado la distribución reportada.",
        "El desbalance de clases se trató desde la evaluación y no mediante reetiquetado artificial del conjunto de prueba. La prevalencia de 13.9% implica que métricas como exactitud pueden ser engañosas: un clasificador sesgado hacia la clase negativa puede parecer competitivo aunque no detecte suficientes casos de riesgo.",
        "El preprocesamiento fue deliberadamente simple. La regresión logística representa un baseline interpretable, mientras que los modelos de ensamble capturan interacciones no lineales entre edad, índice de masa corporal, hipertensión, colesterol, salud general y movilidad. Esta comparación permite aislar el efecto de complejidad del modelo frente a la calidad de variables.",
    ]:
        add_p(doc, text)

    add_heading(doc, "4.2 Métricas y umbrales clínicos", level=2)
    for text in [
        "ROC-AUC se utilizó para medir discriminación global, mientras que PR-AUC fue incorporada por la naturaleza desbalanceada del problema. MCC se reportó como métrica robusta para clasificación binaria porque penaliza simultáneamente falsos positivos y falsos negativos. Brier score y ECE cuantificaron la calidad de las probabilidades.",
        "Los umbrales de decisión no deben fijarse únicamente por maximización matemática. En tamizaje clínico, un umbral sensible puede ser preferible para no omitir pacientes de alto riesgo, siempre que la interfaz explique que el resultado requiere confirmación médica. Por ello, las bandas de riesgo son más útiles que una salida binaria.",
    ]:
        add_p(doc, text)

    add_heading(doc, "4.3 Prevención de fuga de información", level=2)
    for text in [
        "Se evitó usar variables que representaran directamente el diagnóstico o mediciones posteriores al evento objetivo. La separación de entrenamiento, validación y prueba reduce el riesgo de ajustar decisiones sobre el mismo conjunto con el que se reporta desempeño final.",
        "La reproducibilidad se fortaleció al documentar particiones, métricas, figuras y tablas generadas. Esta práctica es importante en modelos clínicos porque pequeñas decisiones de preparación pueden producir diferencias relevantes en AUC, calibración o sensibilidad.",
    ]:
        add_p(doc, text)

    roc = ROOT / "ml-research" / "brfss" / "results" / "models" / "roc_auc_bar.png"
    if roc.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(roc), width=Inches(3.0))
        add_caption(doc, "Figura 1.", "Comparación de ROC-AUC entre modelos evaluados. Fuente: elaboración propia con BRFSS 2015.")

    add_heading(doc, "5. Resultados experimentales")
    add_p(
        doc,
        "Los modelos basados en gradiente lograron el mejor ROC-AUC, pero la diferencia frente a regresión logística fue pequeña. HistGradientBoosting y XGBoost alcanzaron 0.8268, mientras que regresión logística obtuvo 0.8197. Esta brecha de 0.007 sugiere que la ganancia práctica proviene menos de complejidad algorítmica y más de contar con mejores variables clínicas.",
    )
    add_caption(doc, "Tabla 3.", "Comparación resumida de modelos en conjunto de prueba.")
    add_table(doc, ["Modelo", "ROC-AUC", "PR-AUC", "MCC", "Brier"], load_model_rows(), font_size=7.3)

    curve = ROOT / "ml-research" / "brfss" / "results" / "models" / "roc_all.png"
    if curve.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(curve), width=Inches(3.0))
        add_caption(doc, "Figura 2.", "Curvas ROC del benchmark. Fuente: elaboración propia.")

    add_heading(doc, "5.1 Calibración")
    add_p(
        doc,
        "La calibración fue evaluada porque en contextos clínicos una probabilidad mal calibrada puede inducir decisiones equivocadas aun si el ordenamiento por riesgo es aceptable. La calibración isotónica mantuvo Brier score competitivo y redujo ECE en prueba a 0.0029.",
    )
    add_caption(doc, "Tabla 4.", "Calibración del mejor modelo.")
    add_table(
        doc,
        ["Método", "Brier val.", "ECE val.", "Brier test", "ECE test"],
        [
            ["Original", "0.0968", "0.0040", "0.0974", "0.0044"],
            ["Platt", "0.0985", "0.0257", "0.0995", "0.0285"],
            ["Isotónica", "0.0965", "0.0000", "0.0975", "0.0029"],
        ],
        font_size=7.8,
    )
    calib = ROOT / "ml-research" / "brfss" / "results" / "models" / "calib_xgboost.png"
    if calib.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(calib), width=Inches(2.75))
        add_caption(doc, "Figura 3.", "Curva de calibración de XGBoost. Fuente: elaboración propia.")

    add_heading(doc, "5.2 Bandas de riesgo")
    add_p(
        doc,
        "Las bandas de riesgo convierten probabilidades en categorías accionables. El grupo de riesgo bajo presentó 2.9% de prevalencia observada, mientras que el grupo muy alto alcanzó 48.0%. La razón de momios del grupo muy alto fue 30.7 frente al grupo bajo.",
    )
    add_caption(doc, "Tabla 5.", "Prevalencia observada por banda de riesgo.")
    add_table(
        doc,
        ["Banda", "n", "Prev.", "OR", "RR"],
        [
            ["Bajo", "26,149", "2.9%", "1.0", "1.0"],
            ["Moderado", "11,877", "14.2%", "5.5 [5.0, 6.0]", "4.9"],
            ["Alto", "7,844", "29.1%", "13.7 [12.5, 14.9]", "10.0"],
            ["Muy alto", "4,866", "48.0%", "30.7 [28.0, 33.6]", "16.4"],
        ],
        font_size=7.7,
    )

    add_heading(doc, "5.3 Lectura interpretativa", level=2)
    for text in [
        "La lectura de los resultados coincide con conocimiento clínico básico: edad, índice de masa corporal, hipertensión, colesterol alto, salud general percibida y dificultad para caminar concentran información predictiva. Esto no significa causalidad individual, pero sí ofrece señales comunicables para orientar educación y seguimiento.",
        "La explicabilidad cumple dos funciones. Primero, permite detectar si el modelo aprende patrones clínicamente plausibles o artefactos de los datos. Segundo, ayuda al médico a explicar por qué una persona aparece en una banda de riesgo y qué factores podrían atenderse con intervención preventiva.",
        "En un despliegue real, cada predicción debería acompañarse de fecha, versión del modelo, variables disponibles, factores principales y advertencia de uso. La trazabilidad es especialmente importante cuando el sistema se conecta con expediente clínico o genera recomendaciones dentro de una app móvil.",
    ]:
        add_p(doc, text)

    pr_curve = ROOT / "ml-research" / "brfss" / "results" / "models" / "pr_xgboost.png"
    if pr_curve.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(pr_curve), width=Inches(2.9))
        add_caption(doc, "Figura 4.", "Curva precisión-recobrado para XGBoost. Fuente: elaboración propia.")

    add_heading(doc, "5.4 Implicaciones para PREDIA", level=2)
    for text in [
        "El benchmark aporta una base para la capa predictiva de PREDIA, pero su valor aumenta al integrarse con el producto: dashboard clínico, agenda, expediente, signos vitales, alertas y app móvil del paciente. La predicción debe convertirse en tareas concretas, no quedarse como número aislado.",
        "En el perfil médico, la banda de riesgo puede priorizar citas, sugerir revisión de signos vitales y mostrar factores explicativos. En el perfil paciente, la misma información debe expresarse con lenguaje no alarmista, educación preventiva y recordatorios de seguimiento. Esta separación de experiencias reduce riesgo de interpretación incorrecta.",
    ]:
        add_p(doc, text)

    add_heading(doc, "6. Discusión")
    for text in [
        "El resultado central es deliberadamente sobrio: el aprendizaje automático mejora el ordenamiento del riesgo, pero la magnitud de la mejora frente a regresión logística es limitada cuando las variables son autorreportadas y no incluyen biomarcadores. Esto contradice una narrativa común donde modelos más complejos se presentan como solución suficiente.",
        "Para un sistema clínico móvil o web, la utilidad se encuentra en integrar el modelo con acciones: priorizar agenda, activar alertas, sugerir seguimiento, generar explicación de factores de riesgo y registrar la decisión clínica. La predicción aislada tiene menos valor que el circuito completo de atención.",
        "El desempeño también debe evaluarse por seguridad: un modelo no calibrado o presentado como diagnóstico puede aumentar falsos positivos, ansiedad del paciente o retraso de pruebas confirmatorias. La interfaz debe comunicar que el riesgo es una estimación y requiere valoración médica.",
        "Otra implicación es que el despliegue debe distinguir entre exactitud estadística y utilidad clínica. Un AUC aceptable no garantiza que el sistema mejore resultados en salud; para ello se requiere evaluar tiempos de atención, adherencia, seguimiento, aceptación del personal médico y comprensión del paciente.",
        "La moderada diferencia entre modelos también favorece arquitecturas de decisión híbridas. Un modelo simple puede ser suficiente para explicar factores generales, mientras que un ensamble calibrado puede ordenar prioridades operativas. Ambos enfoques pueden coexistir si el sistema conserva trazabilidad, monitoreo y auditoría.",
    ]:
        add_p(doc, text)

    add_note_box(
        doc,
        "Criterio de uso responsable",
        "PREDIA debe presentar el resultado como riesgo estimado, no como diagnóstico. La decisión clínica final corresponde al profesional de salud y debe quedar registrada en expediente.",
    )

    add_heading(doc, "7. Lecciones aprendidas para ML clínico")
    for text in [
        "Primero, los baselines fuertes son obligatorios. Comparar un ensamble contra un método trivial puede inflar conclusiones. Segundo, la calibración no es opcional cuando el resultado se muestra como probabilidad. Tercero, la interpretación debe estar conectada con acciones clínicas, no solo con gráficos técnicos.",
        "Cuarto, el diseño del producto debe evitar automatización excesiva. En PREDIA, el modelo debe actuar como soporte de decisión, con validación médica y registro auditable. Quinto, la mejora futura dependerá de integrar variables de laboratorio, seguimiento longitudinal y datos capturados por paciente desde móvil.",
        "Sexto, la evaluación debe incluir fallos de producto: formularios incompletos, datos fuera de rango, citas duplicadas, pérdida de conexión y perfiles con permisos distintos. En salud digital, la seguridad del modelo depende también de la calidad del flujo de captura y persistencia de datos.",
        "Séptimo, una app móvil aporta utilidad real cuando reduce fricción en tareas concretas: registrar signos, consultar expediente, administrar agenda, iniciar o cerrar atención, descargar documentos e involucrar al paciente en su seguimiento. Esto diferencia una integración móvil funcional de una copia limitada de la web.",
    ]:
        add_p(doc, text)

    add_heading(doc, "8. Limitaciones y amenazas a la validez")
    for text in [
        "La base BRFSS es transversal y autodeclarada, por lo que no permite inferir causalidad ni reemplazar medición clínica. Existen posibles sesgos de recuerdo, subregistro y diferencias de acceso al diagnóstico. La población corresponde a Estados Unidos y puede no representar directamente a otras regiones.",
        "El conjunto no incluye glucosa, hemoglobina glucosilada ni resultados de laboratorio. Tampoco se dispone de temporalidad clínica suficiente para evaluar progresión individual. Aunque se preservaron duplicados para reproducir la base, su existencia puede reflejar patrones de respuesta repetidos y afectar estimaciones si se usa sin control.",
        "La evaluación no sustituye validación prospectiva. Antes de usar el modelo como parte de una decisión asistencial formal, debe probarse con datos locales, revisar desempeño por subgrupos y medir si la intervención reduce tiempos de detección o mejora adherencia al seguimiento.",
        "También existe amenaza de desplazamiento de distribución. Cambios en prevalencia, hábitos, criterios diagnósticos o forma de captura pueden alterar la calibración. Por ello, un sistema productivo debe monitorear desempeño, errores y calidad de datos después del despliegue.",
    ]:
        add_p(doc, text)

    add_heading(doc, "9. Trabajo futuro")
    add_p(
        doc,
        "Como continuación se propone integrar variables clínicas longitudinales, biomarcadores, adherencia terapéutica, signos vitales y datos de autogestión móvil. En PREDIA, esto permitiría actualizar riesgo durante la consulta, registrar cambios desde la app del paciente y validar predicciones contra desenlaces reales.",
    )
    add_p(
        doc,
        "También se recomienda evaluar equidad por edad, sexo, grupo socioeconómico y acceso a servicios de salud, además de construir pruebas de robustez contra fuga de información, drift de datos y cambios de prevalencia.",
    )
    add_p(
        doc,
        "La siguiente etapa técnica debe incluir MLOps básico: versionado de datasets, versionado de modelos, pruebas automáticas de métricas mínimas, monitoreo de calibración y bitácora de predicciones. Esto permitiría detectar degradación antes de que afecte decisiones clínicas.",
    )
    add_p(
        doc,
        "Desde la perspectiva de experiencia móvil, se recomienda validar el modelo con usuarios reales mediante escenarios de consulta: paciente con signos actualizados, paciente con cita vencida, paciente con riesgo alto y paciente con datos incompletos. Cada escenario debe medir comprensión, tiempo de tarea y errores de captura.",
    )

    add_heading(doc, "10. Consideraciones de implementación clínica")
    for text in [
        "La integración del modelo en PREDIA requiere una arquitectura que separe captura de datos, cálculo de riesgo, persistencia clínica y presentación visual. Esta separación permite auditar qué datos entraron al modelo, qué versión produjo la predicción y qué acción tomó el médico después de verla.",
        "En la app móvil, las validaciones deben ejecutarse antes de enviar información a la base de datos: rangos fisiológicos, fechas de cita, campos obligatorios, identidad del paciente y permisos por perfil. Validar en frontend mejora experiencia, pero el backend debe repetir las reglas para evitar registros inválidos por llamadas directas a la API.",
        "La agenda clínica es un punto crítico porque conecta predicción con operación. Un paciente en riesgo alto puede requerir cita prioritaria, pero el sistema debe evitar conflictos de horario, citas en fechas pasadas y cambios sin trazabilidad. Iniciar, finalizar, cancelar y reagendar deben quedar reflejados tanto en móvil como en web.",
        "La interfaz debe usar microcopys de seguridad: 'riesgo estimado', 'requiere valoración', 'datos incompletos' y 'última actualización'. Estos textos reducen la posibilidad de que el paciente interprete una probabilidad como diagnóstico definitivo. Para médicos, el panel debe mostrar factores principales y estado de calibración, no solo un color de riesgo.",
        "Finalmente, el despliegue debe contemplar monitoreo operativo: latencia de API, errores de autenticación, fallos al guardar expedientes, generación de PDF, sincronización móvil-web y eventos de seguridad. Un modelo útil puede perder valor si el sistema falla en captura, persistencia o comunicación clínica.",
    ]:
        add_p(doc, text)

    add_heading(doc, "11. Gobernanza, privacidad y seguridad")
    for text in [
        "Un sistema clínico inteligente debe diseñarse bajo el principio de mínimo privilegio. El paciente no debe acceder a expedientes ajenos y el médico solo debe consultar o modificar datos dentro de su ámbito operativo. La protección mediante autenticación, tokens JWT, roles y validación de permisos es parte de la seguridad del modelo, no un elemento separado.",
        "La privacidad también depende de reducir exposición innecesaria. Las respuestas de API deben devolver únicamente los campos requeridos por la pantalla activa, evitando incluir datos sensibles que no se muestran. En móvil, la sesión debe almacenarse de forma segura y los errores no deben revelar detalles internos de base de datos o infraestructura.",
        "Toda predicción relevante debería almacenarse con metadatos: identificador de paciente, fecha, versión de modelo, variables disponibles, probabilidad, banda de riesgo, factores explicativos y usuario que la consultó. Esta bitácora permite auditoría posterior y facilita explicar por qué una decisión fue tomada.",
        "La gobernanza operacional requiere monitorear calidad de datos. Valores imposibles, campos incompletos, fechas inválidas o cambios repetidos en registros clínicos pueden degradar el desempeño de forma silenciosa. Las validaciones de formularios, logs y alertas operativas son mecanismos preventivos contra ese deterioro.",
        "En etapa académica, la demostración debe mostrar que las validaciones rechazan entradas fuera de rango y que las acciones móviles se reflejan en la web. En etapa productiva, el mismo principio debe ampliarse con pruebas automatizadas, monitoreo de errores, control de versiones y revisión periódica por especialistas clínicos.",
    ]:
        add_p(doc, text)

    add_heading(doc, "12. Conclusiones")
    for text in [
        "El benchmark muestra que el aprendizaje automático aporta una mejora real pero moderada para predecir riesgo de diabetes con indicadores poblacionales. Los mejores modelos alcanzaron ROC-AUC cercano a 0.827, mientras que regresión logística se mantuvo competitiva con 0.820.",
        "La conclusión práctica es que el valor clínico no depende solo del algoritmo. Depende de calibración, interpretabilidad, validación externa, integración al flujo de atención y comunicación responsable con el usuario. El modelo puede apoyar cribado y priorización, pero no debe emplearse como diagnóstico autónomo.",
    ]:
        add_p(doc, text)

    add_heading(doc, "Declaración de reproducibilidad")
    add_p(
        doc,
        "El estudio declara uso de BRFSS 2015, partición 60/20/20, comparación de nueve estimadores y generación de métricas y figuras reproducibles. Los archivos del proyecto PREDIA incluyen tablas CSV de resultados, curvas ROC/PR, calibración y matrices de confusión bajo ml-research/brfss/results/models.",
        size=9.5,
    )
    add_p(
        doc,
        "Para repetir el benchmark se requieren tres elementos: acceso al dataset público, entorno de Python con bibliotecas de aprendizaje automático y ejecución del pipeline de evaluación. La salida esperada incluye comparativas de modelos, curvas, métricas de calibración y tablas de bandas de riesgo.",
        size=9.5,
    )

    add_heading(doc, "Declaraciones")
    add_p(doc, "Conflicto de interés: no se encontró una declaración formal en el material fuente; requiere confirmación por los autores antes del envío.", size=9.2)
    add_p(doc, "Financiamiento: no se identificó financiamiento externo específico en el material fuente.", size=9.2)
    add_p(doc, "Disponibilidad de datos: BRFSS 2015 es un conjunto público del CDC; los resultados procesados se documentan en los archivos del proyecto.", size=9.2)
    add_p(doc, "Contribución de autores: pendiente de desglosar por autor en la versión final de envío.", size=9.2)

    add_heading(doc, "Referencias")
    if refs:
        for ref in refs:
            add_p(doc, ref, size=7.5, after=1)
    else:
        add_p(doc, "Referencias tomadas del PDF fuente. Se recomienda verificar DOI/URL antes del envío.", size=8)

    FINAL_DOCX.unlink(missing_ok=True)
    doc.save(FINAL_DOCX)


if __name__ == "__main__":
    build_document(sys.argv[1] if len(sys.argv) > 1 else "1-14")
    print(FINAL_DOCX)
