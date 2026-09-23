#!/usr/bin/env python3
from __future__ import annotations

import csv
import math
import re
import shutil
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path("/home/wowo/Descargas/cimerti-diabetes.docx")
OUTPUT = ROOT / "docs" / "cimerti-diabetes-corregido.docx"
BBL = ROOT / "paper" / "paper_ml_diabetes_mimeti2026_es.bbl"
RESULTS = ROOT / "ml-research" / "brfss" / "results" / "models"
TABLE_HELPERS = Path(
    "/home/wowo/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.826.12353/skills/documents/scripts"
)
sys.path.insert(0, str(TABLE_HELPERS))
from table_geometry import apply_table_geometry  # noqa: E402


TNR = "Times New Roman"
LINK_BLUE = RGBColor(5, 99, 193)
INK = "111111"
NAVY = "17365D"
RULE = "7F7F7F"
LIGHT = "E7EDF4"
COL_WIDTH_DXA = 4500


def font_path(bold: bool = False) -> str:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    return next(path for path in candidates if Path(path).exists())


def chart_font(size: int, bold: bool = False):
    return ImageFont.truetype(font_path(bold), size)


def clean_latex_reference(text: str) -> str:
    text = re.sub(r"%.*", "", text)
    text = text.replace("~", " ").replace("---", " - ")
    text = text.replace("``", '"').replace("''", '"')
    text = text.replace("\\newblock", " ")
    text = re.sub(r"\\BIBentry[A-Za-z]+", " ", text)
    text = re.sub(r"\\hskip.*?\\relax", " ", text)
    text = re.sub(r"\\url\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\emph\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\BIBforeignlanguage\{[^{}]+\}\{([^{}]+)\}", r"\1", text)
    text = text.replace("{\\'E}", "É").replace("{\\v{s}}", "š")
    text = re.sub(r"\\[a-zA-Z]+(?:\[[^\]]+\])?", " ", text)
    text = text.replace("{", "").replace("}", "").replace("\\", "")
    text = text.replace("--", "-")
    return re.sub(r"\s+", " ", text).strip(" ,.;")


def parse_references() -> list[str]:
    text = BBL.read_text(encoding="utf-8", errors="ignore")
    refs: list[str] = []
    for idx, chunk in enumerate(text.split("\\bibitem")[1:], start=1):
        chunk = re.sub(r"^\{[^{}]+\}", "", chunk, count=1)
        chunk = chunk.split("\\end{thebibliography}")[0]
        cleaned = clean_latex_reference(chunk)
        if cleaned:
            refs.append(f"[{idx}] {cleaned}.")
    if len(refs) != 35:
        raise RuntimeError(f"Expected 35 references, found {len(refs)}")
    return refs


def load_model_rows() -> list[list[str]]:
    rows: list[list[str]] = []
    with (RESULTS / "comparison_test.csv").open(newline="", encoding="utf-8") as stream:
        for row in csv.DictReader(stream):
            name = row["model"].replace("_", " ").title()
            name = name.replace("Hist Gradient Boosting", "Hist. Gradient Boost.")
            name = (
                name.replace("Mlp", "MLP")
                .replace("Knn", "KNN")
                .replace("Svm", "SVM")
                .replace("Xgboost", "XGBoost")
                .replace("Lightgbm", "LightGBM")
            )
            ci = row["ci"].strip("[]").replace(",", "-")
            rows.append(
                [
                    name,
                    f"{float(row['roc_auc']):.3f}",
                    ci,
                    f"{float(row['pr_auc']):.3f}",
                    f"{float(row['mcc']):.3f}",
                    f"{float(row['brier']):.3f}",
                ]
            )
    return rows


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def format_run(run, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = TNR
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), TNR)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    size: float,
    bold: bool = False,
    align=None,
) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    format_run(run, size, bold=bold)
    if align is not None:
        paragraph.alignment = align


def add_hyperlink(paragraph, text: str, url: str, *, size: float = 10.0) -> None:
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), TNR)
    fonts.set(qn("w:hAnsi"), TNR)
    fonts.set(qn("w:eastAsia"), TNR)
    run_props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    run_props.append(size_el)
    run.append(run_props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def wrap_run_with_hyperlink(paragraph, run, url: str) -> None:
    parent = run._r.getparent()
    if parent is None or parent.tag != qn("w:p"):
        return
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    index = parent.index(run._r)
    parent.remove(run._r)
    run.font.color.rgb = LINK_BLUE
    run.font.underline = True
    hyperlink.append(run._r)
    parent.insert(index, hyperlink)


def link_affiliation(paragraph, url: str) -> None:
    runs = list(paragraph.runs)
    start = next((i for i, run in enumerate(runs) if "Universidad" in run.text), None)
    if start is None:
        return
    for run in runs[start:]:
        if run._r.xpath(".//w:drawing"):
            break
        wrap_run_with_hyperlink(paragraph, run, url)


def link_matching_run(paragraph, token: str, url: str) -> None:
    for run in list(paragraph.runs):
        if token in run.text:
            wrap_run_with_hyperlink(paragraph, run, url)
            return


def patch_first_page(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    set_paragraph_text(
        paragraphs[0],
        "How Much Does Machine Learning Really Improve Diabetes Risk Prediction from "
        "Population Health Indicators? An Honest, Interpretable, and Transparent "
        "Benchmark on BRFSS 2015",
        size=14,
        bold=True,
    )
    set_paragraph_text(
        paragraphs[2],
        "¿Cuánto Mejora Realmente el Aprendizaje Automático la Predicción del Riesgo de "
        "Diabetes a partir de Indicadores de Salud Poblacional? Un Benchmark Honesto, "
        "Interpretable y Transparente sobre BRFSS 2015",
        size=14,
        bold=True,
    )

    byline = paragraphs[4]
    clear_paragraph(byline)
    author_parts = [
        ("Álvarez-Sánchez, Guillermo", "a", False),
        (", Peredo-Valderrama, Iván", "b", True),
        (", Moya-Moya, José Javier", "c", False),
        (" and Pacindo-Piña, Jairo", "d", False),
    ]
    for name, marker, corresponding in author_parts:
        run = byline.add_run(name)
        format_run(run, 14)
        mark = byline.add_run(marker + ("*" if corresponding else ""))
        format_run(mark, 10)
        mark.font.superscript = True
    byline.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for run in paragraphs[6].runs:
        if run.text == "SNP":
            run.text = "CONAHCYT"
            format_run(run, 9.5)
    for run in paragraphs[7].runs:
        if run.text == "SNI- SECIHTI":
            run.text = "SNI-SECIHTI"
            format_run(run, 9.5)

    ror = "https://ror.org/0580rm578"
    author_links = [
        (paragraphs[6], "QXJ-0529-2026", "0009-0005-6254-143X", "2044833"),
        (paragraphs[7], "QXI-7175-2026", "0000-0002-2484-9557", "167230"),
        (paragraphs[8], "QXJ-0392-2026", "0009-0007-3628-5000", "1292895"),
        (paragraphs[9], "QXJ-0480-2026", "0009-0008-4895-378X", "2089135"),
    ]
    for paragraph, clarivate, orcid, secihti in author_links:
        link_affiliation(paragraph, ror)
        link_matching_run(
            paragraph,
            clarivate,
            f"https://www.webofscience.com/wos/author/record/{clarivate}",
        )
        link_matching_run(paragraph, orcid, f"https://orcid.org/{orcid}")
        link_matching_run(paragraph, secihti, "https://secihti.mx/")

    remove_paragraph(paragraphs[11])
    remove_paragraph(paragraphs[12])
    remove_paragraph(paragraphs[13])

    classification = paragraphs[16]
    clear_paragraph(classification)
    classification.add_run("Classification reference: ")
    for run in classification.runs:
        format_run(run, 9.5)
    add_hyperlink(classification, "MARVID research areas", "https://marvid.org/research_areas.php", size=9.5)

    set_paragraph_text(
        paragraphs[21],
        "DOI: Pending assignment by ECORFAN",
        size=9.5,
        bold=True,
    )

    citation = (
        "Citation: Álvarez-Sánchez, Guillermo, Peredo-Valderrama, Iván, Moya-Moya, José "
        "Javier, Pacindo-Piña, Jairo. How Much Does Machine Learning Really Improve "
        "Diabetes Risk Prediction from Population Health Indicators? An Honest, "
        "Interpretable, and Transparent Benchmark on BRFSS 2015. ECORFAN Journal-Mexico. "
        "2026. V-N: Pages."
    )
    set_paragraph_text(paragraphs[44], citation, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    correspondence = paragraphs[46]
    clear_paragraph(correspondence)
    run = correspondence.add_run("* Corresponding author: ")
    format_run(run, 10)
    add_hyperlink(correspondence, "ivan.peredo@upq.mx", "mailto:ivan.peredo@upq.mx", size=10)

    key_cell = doc.tables[0].cell(1, 0)
    key_cell.text = ""
    key_p = key_cell.paragraphs[0]
    key_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    key_run = key_p.add_run(
        "Provides a transparent comparison of nine diabetes risk models and demonstrates "
        "that calibration, interpretability, and admissible variables matter more than "
        "marginal gains in ROC-AUC."
    )
    format_run(key_run, 10)


def patch_footer_package(path: Path) -> None:
    """Replace editorial footer instructions without resizing its floating layout."""
    citation = (
        "Álvarez-Sánchez, Guillermo, Peredo-Valderrama, Iván, Moya-Moya, José Javier "
        "and Pacindo-Piña, Jairo. How Much Does Machine Learning Really Improve Diabetes "
        "Risk Prediction from Population Health Indicators? An Honest, Interpretable, "
        "and Transparent Benchmark on BRFSS 2015. ECORFAN Journal-Mexico. "
        "2026. V-N: Pages. DOI: Pending assignment by ECORFAN."
    )
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path, "r") as source:
        entries = {item.filename: source.read(item.filename) for item in source.infolist()}

    footer_name = "word/footer1.xml"
    root = etree.fromstring(entries[footer_name])
    text_nodes = root.xpath(".//w:t", namespaces=namespace)
    for start in (0, 11):
        text_nodes[start].text = citation
        for index in range(start + 1, start + 11):
            text_nodes[index].text = ""
    entries[footer_name] = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )

    patched = path.with_suffix(".patched.docx")
    with ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)
    patched.replace(path)


def extract_template_images(work_dir: Path) -> dict[str, Path]:
    mapping = {
        "graphical_en": "word/media/image12.png",
        "graphical_es": "word/media/image13.png",
        "shap": "word/media/image17.png",
        "permutation": "word/media/image18.png",
    }
    extracted: dict[str, Path] = {}
    with ZipFile(REFERENCE) as archive:
        for key, member in mapping.items():
            target = work_dir / Path(member).name
            target.write_bytes(archive.read(member))
            extracted[key] = target
    patch_graphical_contribution(extracted["graphical_en"], language="en")
    patch_graphical_contribution(extracted["graphical_es"], language="es")
    return extracted


def patch_graphical_contribution(path: Path, *, language: str) -> None:
    """Keep the supplied graphical abstract while correcting its contribution claim."""
    image = Image.open(path).convert("RGB")
    draw = ImageDraw.Draw(image)
    serif = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"

    if language == "en":
        draw.rectangle((145, 731, 699, 778), fill="white")
        draw.text((155, 735), "•", fill="#000000", font=ImageFont.truetype(serif, 18))
        font = ImageFont.truetype(serif, 14)
        draw.text(
            (178, 735),
            "Provides an honest and transparent benchmark of ML models for",
            fill="#000000",
            font=font,
        )
        draw.text(
            (178, 754),
            "diabetes risk prediction using population health indicators.",
            fill="#000000",
            font=font,
        )
    else:
        draw.rectangle((96, 507, 477, 538), fill="white")
        draw.text((103, 510), "•", fill="#000000", font=ImageFont.truetype(serif, 11))
        font = ImageFont.truetype(serif, 8)
        draw.text(
            (118, 510),
            "Proporciona un referente honesto y transparente de modelos de aprendizaje automático",
            fill="#000000",
            font=font,
        )
        draw.text(
            (118, 522),
            "para la predicción del riesgo de diabetes utilizando indicadores de salud poblacional.",
            fill="#000000",
            font=font,
        )

    image.save(path)


def make_auc_chart(path: Path, rows: list[list[str]]) -> None:
    width, height = 1180, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = chart_font(38, True)
    label = chart_font(24)
    value = chart_font(22, True)
    draw.text((60, 35), "Held-out test ROC-AUC by model", font=title, fill="#17365D")
    left, top, right, bottom = 340, 120, 1080, 690
    draw.line((left, top, left, bottom), fill="#666666", width=2)
    draw.line((left, bottom, right, bottom), fill="#666666", width=2)
    min_auc, max_auc = 0.65, 0.84
    bar_h, gap = 42, 18
    for index, row in enumerate(rows):
        model, auc = row[0], float(row[1])
        y = top + index * (bar_h + gap)
        draw.text((55, y + 5), model, font=label, fill="#222222")
        x0 = left + int((min_auc - 0.60) / (max_auc - 0.60) * (right - left))
        x1 = left + int((auc - 0.60) / (max_auc - 0.60) * (right - left))
        color = "#1F4E79" if index < 3 else "#6E879F"
        if model == "Logistic Regression":
            color = "#B45F06"
        draw.rounded_rectangle((x0, y, x1, y + bar_h), radius=6, fill=color)
        draw.text((x1 + 10, y + 6), f"{auc:.3f}", font=value, fill="#222222")
    for tick in (0.65, 0.70, 0.75, 0.80, 0.84):
        x = left + int((tick - 0.60) / (max_auc - 0.60) * (right - left))
        draw.line((x, bottom, x, bottom + 8), fill="#666666", width=2)
        draw.text((x - 25, bottom + 12), f"{tick:.2f}", font=label, fill="#444444")
    image.save(path)


def make_calibration_chart(path: Path) -> None:
    width, height = 1180, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = chart_font(38, True)
    label = chart_font(24)
    small = chart_font(20)
    bold = chart_font(22, True)
    draw.text((60, 35), "Probability calibration on the held-out test set", font=title, fill="#17365D")
    variants = ["Original", "Platt", "Isotonic"]
    brier = [0.0974, 0.0995, 0.0975]
    ece = [0.0044, 0.0285, 0.0029]
    panels = [(80, 125, 545, 620, "Brier score", brier, 0.105), (635, 125, 1100, 620, "Expected calibration error", ece, 0.032)]
    colors = ["#4C78A8", "#E08B36", "#2A9D8F"]
    for x0, y0, x1, y1, panel_title, values, vmax in panels:
        draw.text((x0, y0 - 5), panel_title, font=bold, fill="#222222")
        axis_top = y0 + 55
        draw.line((x0 + 55, axis_top, x0 + 55, y1), fill="#777777", width=2)
        draw.line((x0 + 55, y1, x1, y1), fill="#777777", width=2)
        bar_w, gap = 90, 45
        for i, (name, metric) in enumerate(zip(variants, values)):
            bx0 = x0 + 85 + i * (bar_w + gap)
            bh = int(metric / vmax * (y1 - axis_top - 35))
            by0 = y1 - bh
            draw.rounded_rectangle((bx0, by0, bx0 + bar_w, y1), radius=6, fill=colors[i])
            draw.text((bx0 - 5, y1 + 12), name, font=small, fill="#333333")
            draw.text((bx0 - 2, by0 - 30), f"{metric:.4f}", font=small, fill="#222222")
    draw.text((60, 670), "Lower values indicate better probabilistic calibration.", font=label, fill="#444444")
    image.save(path)


def make_risk_or_chart(path: Path) -> None:
    width, height = 1100, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = chart_font(38, True)
    label = chart_font(26)
    small = chart_font(22)
    draw.text((60, 35), "Odds ratios by calibrated risk band", font=title, fill="#17365D")
    bands = ["Low", "Moderate", "High", "Very high"]
    odds = [1.0, 5.5, 13.7, 30.7]
    lows = [1.0, 5.0, 12.5, 28.0]
    highs = [1.0, 6.0, 14.9, 33.6]
    left, right, top, bottom = 250, 1010, 120, 610
    log_min, log_max = 0.0, math.log10(40)
    def xpos(value: float) -> int:
        return left + int((math.log10(value) - log_min) / (log_max - log_min) * (right - left))
    draw.line((left, top, left, bottom), fill="#777777", width=2)
    draw.line((left, bottom, right, bottom), fill="#777777", width=2)
    draw.line((xpos(1), top, xpos(1), bottom), fill="#999999", width=2)
    for i, band in enumerate(bands):
        y = top + 65 + i * 115
        draw.text((55, y - 18), band, font=label, fill="#222222")
        draw.line((xpos(lows[i]), y, xpos(highs[i]), y), fill="#1F4E79", width=7)
        draw.ellipse((xpos(odds[i]) - 12, y - 12, xpos(odds[i]) + 12, y + 12), fill="#B45F06")
        draw.text((xpos(highs[i]) + 14, y - 16), f"OR {odds[i]:.1f}", font=small, fill="#333333")
    for tick in (1, 2, 5, 10, 20, 40):
        x = xpos(tick)
        draw.line((x, bottom, x, bottom + 8), fill="#666666", width=2)
        draw.text((x - 14, bottom + 15), str(tick), font=small, fill="#444444")
    draw.text((420, 670), "Logarithmic odds-ratio scale", font=label, fill="#444444")
    image.save(path)


def make_risk_distribution_chart(path: Path) -> None:
    width, height = 1100, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = chart_font(34, True)
    label = chart_font(26)
    small = chart_font(22)
    bold = chart_font(23, True)
    draw.text(
        (width // 2, 35),
        "Population distribution and observed prevalence",
        font=title,
        fill="#17365D",
        anchor="ma",
    )
    bands = ["Low", "Moderate", "High", "Very high"]
    counts = [26149, 11877, 7844, 4866]
    prevalence = [2.9, 14.2, 29.1, 48.0]
    colors = ["#557A74", "#D7B56D", "#D8874B", "#A33D3D"]
    left, right, top, bottom = 120, 1020, 135, 610
    bar_w, gap = 150, 65
    max_count = 28000
    draw.line((left, bottom, right, bottom), fill="#777777", width=2)
    for i, (band, count, prev) in enumerate(zip(bands, counts, prevalence)):
        x0 = left + 55 + i * (bar_w + gap)
        bh = int(count / max_count * (bottom - top - 30))
        y0 = bottom - bh
        draw.rounded_rectangle((x0, y0, x0 + bar_w, bottom), radius=8, fill=colors[i])
        draw.text((x0 + 14, y0 - 35), f"n={count:,}", font=small, fill="#333333")
        draw.text((x0 + 30, y0 + 20), f"{prev:.1f}%", font=bold, fill="white")
        draw.text((x0 + 12, bottom + 15), band, font=label, fill="#333333")
    draw.text((60, 670), "Percentages inside bars are observed diabetes prevalence.", font=label, fill="#444444")
    image.save(path)


def translate_permutation_chart(source: Path, target: Path) -> None:
    image = Image.open(source).convert("RGB")
    draw = ImageDraw.Draw(image)
    labels = [
        "Perceived general health",
        "Body mass index (BMI)",
        "Age group",
        "High blood pressure",
        "High cholesterol",
        "Cholesterol check in 5 years",
        "Heavy alcohol consumption",
        "Sex (0=female, 1=male)",
        "Income level",
        "Heart disease or heart attack",
        "Difficulty walking/climbing stairs",
        "Poor mental-health days (30d)",
        "History of stroke",
        "Education level",
        "Smoker (>=100 lifetime cigarettes)",
        "Poor physical-health days (30d)",
        "Fruit consumption >=1/day",
        "Vegetable consumption >=1/day",
        "Has healthcare coverage",
        "Could not see doctor due to cost",
    ]
    label_font = chart_font(7)
    axis_font = chart_font(9)
    draw.rectangle((0, 24, 163, 347), fill="white")
    for index, text in enumerate(labels):
        draw.text((160, 36 + index * 15), text, font=label_font, fill="#222222", anchor="rm")
    draw.rectangle((292, 366, 380, 382), fill="white")
    draw.text((336, 374), "Importance", font=axis_font, fill="#222222", anchor="mm")
    image.save(target)


def copy_role(paragraph) -> tuple[object | None, object | None]:
    ppr = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    source_run = next((run for run in paragraph.runs if run.text.strip()), None)
    rpr = deepcopy(source_run._r.rPr) if source_run is not None and source_run._r.rPr is not None else None
    return ppr, rpr


def apply_role(paragraph, role: tuple[object | None, object | None]) -> None:
    ppr, _ = role
    existing = paragraph._p.pPr
    if existing is not None:
        paragraph._p.remove(existing)
    if ppr is not None:
        paragraph._p.insert(0, deepcopy(ppr))


def add_role_paragraph(
    doc: Document,
    text: str,
    role: tuple[object | None, object | None],
    *,
    size: float,
    bold: bool = False,
    italic: bool = False,
    after: float = 4,
    keep_with_next: bool = False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
):
    paragraph = doc.add_paragraph()
    apply_role(paragraph, role)
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    _, rpr = role
    if rpr is not None:
        current = run._r.rPr
        if current is not None:
            run._r.remove(current)
        run._r.insert(0, deepcopy(rpr))
    format_run(run, size, bold=bold, italic=italic)
    return paragraph


def set_cell_border(cell, *, top=None, bottom=None, left=None, right=None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        val, color, size = spec
        element.set(qn("w:val"), val)
        element.set(qn("w:color"), color)
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")


def add_scientific_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
    *,
    font_size: float,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Normal"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        format_run(run, font_size, bold=True)
        set_cell_border(cell, top=("single", INK, 8), bottom=("single", INK, 8))
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            format_run(run, font_size)
        row_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=("single", INK, 8))
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, width: float, caption: str, roles) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split(".", 1)[0])
    add_role_paragraph(doc, caption, roles["caption"], size=9.5, after=5)


def strip_scientific_body(doc: Document) -> dict[str, tuple[object | None, object | None]]:
    paragraphs = list(doc.paragraphs)

    def find_paragraph(text: str, start: int = 0):
        for paragraph in paragraphs[start:]:
            if paragraph.text.strip() == text:
                return paragraph
        raise ValueError(f"Required template paragraph not found: {text}")

    abstract_heading = find_paragraph("Abstract")
    abstract_index = paragraphs.index(abstract_heading)
    abstract_body = next(
        paragraph
        for paragraph in paragraphs[abstract_index + 1 :]
        if paragraph.text.strip()
    )
    methodology_subheading = find_paragraph("A. Protocol and Leakage Control")
    reference_heading = find_paragraph("References")
    reference_index = paragraphs.index(reference_heading)
    reference_body = next(
        paragraph
        for paragraph in paragraphs[reference_index + 1 :]
        if paragraph.text.strip()
    )
    caption_paragraph = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text.strip().startswith("Figure")
    )
    roles = {
        "heading": copy_role(abstract_heading),
        "body": copy_role(abstract_body),
        "subheading": copy_role(methodology_subheading),
        "caption": copy_role(caption_paragraph),
        "reference": copy_role(reference_body),
    }
    start = abstract_heading._p
    body = doc._element.body
    start_index = list(body).index(start)
    for child in list(body)[start_index:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return roles


def build_scientific_body(doc: Document, roles, assets: dict[str, Path], work_dir: Path) -> None:
    rows = load_model_rows()
    auc_chart = work_dir / "figure_auc.png"
    calibration_chart = work_dir / "figure_calibration.png"
    risk_or_chart = work_dir / "figure_risk_or.png"
    risk_distribution_chart = work_dir / "figure_risk_distribution.png"
    permutation_chart = work_dir / "figure_permutation_en.png"
    make_auc_chart(auc_chart, rows)
    make_calibration_chart(calibration_chart)
    make_risk_or_chart(risk_or_chart)
    make_risk_distribution_chart(risk_distribution_chart)
    translate_permutation_chart(assets["permutation"], permutation_chart)

    def heading(text: str) -> None:
        add_role_paragraph(doc, text, roles["heading"], size=12, bold=True, after=5, keep_with_next=True)

    def subheading(text: str) -> None:
        add_role_paragraph(doc, text, roles["subheading"], size=12, after=4, keep_with_next=True)

    def body(text: str) -> None:
        add_role_paragraph(doc, text, roles["body"], size=11.5, after=5)

    def caption(text: str) -> None:
        add_role_paragraph(doc, text, roles["caption"], size=9.5, after=3, keep_with_next=True)

    heading("Abstract")
    body(
        "Machine learning (ML) is increasingly used for population-level diabetes risk prediction, often under the premise that it substantially outperforms classical statistical models. This study rigorously evaluates the incremental benefit of different ML algorithms over logistic regression in a pre-laboratory screening setting designed to avoid direct diagnostic-label leakage from laboratory criteria. We used the CDC Diabetes Health Indicators dataset (BRFSS 2015; n = 253,680; prevalence 13.9%), which contains only self-reported survey indicators. Nine estimators were evaluated under a uniform protocol with a stratified 60/20/20 split, randomized hyperparameter search, preprocessing fitted only within training folds, and ROC-AUC assessment on a held-out test set. Uncertainty was estimated with 1,000 bootstrap resamples. Calibration, risk stratification, and interpretability were also examined. Histogram gradient boosting and XGBoost achieved ROC-AUC = 0.827 (95% CI: 0.822-0.832 and 0.822-0.831, respectively), compared with 0.820 (0.815-0.824) for L2-regularized logistic regression. The paired difference was small (ΔAUC = 0.007; bootstrap 95% CI: 0.006-0.009), statistically detectable because of the large sample but of limited practical importance. The leading boosting models were mutually indistinguishable. Seven of the nine estimators clustered between ROC-AUC 0.819 and 0.827, whereas KNN and SVM performed below that range. Accuracy (0.866) was minimally informative relative to the majority-class baseline (0.861). Isotonic calibration achieved an expected calibration error of 0.0029, and four risk levels separated observed prevalence from 2.9% to 48.0% (odds ratio 30.7; 95% CI: 28.0-33.6). Algorithm selection therefore provides only marginal gains when the variables constrain predictive information. Leakage control, calibration, interpretability, and transparent reporting are more important than maximizing small AUC differences. These models may support population screening and triage, but they do not replace clinical diagnosis."
    )
    add_figure(doc, assets["graphical_en"], 3.12, "Graphical abstract. Study objective, leakage-control methodology, and clinical contribution.", roles)
    heading("Keywords")
    body("Diabetes risk prediction; machine learning; BRFSS; calibration; explainable artificial intelligence; information leakage; transparent reporting.")

    heading("Resumen")
    body(
        "El aprendizaje automático se utiliza cada vez más para predecir el riesgo poblacional de diabetes, con frecuencia bajo la premisa de que supera ampliamente a los modelos estadísticos clásicos. Este estudio evalúa rigurosamente el beneficio incremental de distintos algoritmos frente a la regresión logística en un escenario de cribado prelaboratorio diseñado para evitar la fuga directa de la etiqueta diagnóstica desde criterios de laboratorio. Se utilizó el conjunto Diabetes Health Indicators de los CDC (BRFSS 2015; n = 253,680; prevalencia 13.9%), compuesto únicamente por indicadores de encuesta autorreportados. Nueve estimadores fueron evaluados con una partición estratificada 60/20/20, búsqueda aleatorizada de hiperparámetros, preprocesamiento ajustado solo dentro de los pliegues de entrenamiento y evaluación ROC-AUC en una prueba retenida. La incertidumbre se estimó con 1,000 remuestreos bootstrap. Histogram gradient boosting y XGBoost alcanzaron ROC-AUC = 0.827 (IC 95%: 0.822-0.832 y 0.822-0.831, respectivamente), frente a 0.820 (0.815-0.824) de la regresión logística regularizada. La diferencia emparejada fue pequeña (ΔAUC = 0.007; IC bootstrap: 0.006-0.009) y de relevancia práctica limitada. Los principales modelos de boosting fueron mutuamente indistinguibles. Siete de los nueve estimadores se agruparon entre ROC-AUC 0.819 y 0.827, mientras que KNN y SVM rindieron por debajo de ese intervalo. La exactitud de 0.866 fue poco informativa frente a la línea base de la clase mayoritaria de 0.861. La calibración isotónica alcanzó ECE = 0.0029 y cuatro bandas de riesgo separaron la prevalencia observada de 2.9% a 48.0% (razón de momios 30.7; IC 95%: 28.0-33.6). Controlar la fuga, calibrar, interpretar y reportar con transparencia aporta más valor que maximizar diferencias pequeñas de AUC. El modelo apoya el cribado y el triaje, pero no sustituye el diagnóstico clínico."
    )
    add_figure(doc, assets["graphical_es"], 3.12, "Resumen gráfico. Objetivo, metodología con control de fuga y contribución clínica.", roles)
    heading("Palabras clave")
    body("Predicción de riesgo de diabetes; aprendizaje automático; BRFSS; calibración; inteligencia artificial explicable; fuga de información; transparencia metodológica.")

    heading("Introduction")
    body(
        "Type 2 diabetes is one of the most prevalent chronic diseases worldwide, and identifying people at risk before complications appear is a public-health priority [1], [2]. Population health surveys provide low-cost information about demographic, behavioral, and cardiometabolic factors, while modern gradient-boosted trees and neural models offer flexible ways to exploit those variables. This combination has produced extensive literature reporting useful ML performance for diabetes prediction [3]-[5]."
    )
    body(
        "A recurring claim is that nonlinear algorithms substantially outperform logistic regression. Yet a systematic review of clinical prediction studies found no consistent ML advantage when comparisons were conducted fairly [7]. Reported improvements may instead reflect richer variables, optimistic validation, information leakage, or weak baselines. In high-stakes settings, a small gain in discrimination cannot compensate for poor calibration, opaque behavior, or invalid data provenance."
    )
    body(
        "We ask how much modern ML actually improves diabetes prediction when the inputs are ordinary, self-reported population health indicators. Four hypotheses guide the analysis: (H1) complex models will not materially outperform a well-regularized linear baseline; (H2) the available variables impose an observable performance plateau; (H3) calibration and interpretability offer greater clinical value than small gains in ROC-AUC; and (H4) explicit control of information leakage changes the honest interpretation of diabetes models."
    )
    body(
        "BRFSS contains no diagnostic laboratory measurements such as HbA1c or fasting glucose. It therefore represents genuine pre-laboratory screening and avoids the common error of using diagnostic criteria as predictors [8], [9]. The contribution is a uniform benchmark of nine estimators, uncertainty quantified with paired bootstrap resampling, probability calibration, clinically interpretable risk bands, and a transparent evaluation protocol. The goal is screening support, not automated diagnosis."
    )

    heading("Related Work")
    subheading("A. Survey-based diabetes prediction")
    body(
        "Previous studies using population-health and questionnaire-derived predictors generally report useful but imperfect discrimination [3]-[5]. Neural networks and ensembles sometimes lead point estimates, but their gains over simpler models are often modest. Laboratory-rich datasets, including those in [6] and [34], address a different prediction setting and should not be compared directly with pre-laboratory survey screening."
    )
    subheading("B. Machine learning versus regression")
    body(
        "The comparative literature cautions against treating algorithmic complexity as evidence of clinical benefit. Christodoulou et al. found no systematic performance advantage of ML over logistic regression [7], while Steyerberg emphasizes validation, calibration, and transportability over discrimination alone [10]. A fair benchmark must therefore include a strong regularized linear baseline and paired uncertainty estimates."
    )
    subheading("C. Calibration, imbalance, and interpretability")
    body(
        "Reliable probabilities are essential when predictions define triage thresholds [11]-[13]. Because diabetes prevalence is imbalanced, PR-AUC and MCC complement ROC-AUC and raw accuracy [20]-[22]. SHAP, permutation importance, and interpretable models can expose the factors driving risk [14]-[18], but post-hoc explanations should not be treated as causal or automatically trustworthy [19]. These concerns motivate the combined evaluation used here."
    )

    heading("Dataset")
    subheading("A. Source and cohort")
    body(
        "We used the public Diabetes Health Indicators dataset derived from the 2015 Behavioral Risk Factor Surveillance System [23], [24]. The binary file contains 253,680 respondents, 21 predictors, and the target Diabetes_binary. The positive prevalence is 13.9%, yielding a negative-to-positive ratio of 6.18:1. The dataset has no missing values. Exact duplicate response vectors were retained because categorical survey profiles can legitimately recur across different respondents; removing them without survey identifiers would alter prevalence and population structure."
    )
    caption("Table I. Principal characteristics of the BRFSS 2015 binary cohort.")
    add_scientific_table(
        doc,
        ["Property", "Value"],
        [
            ["Respondents", "253,680"],
            ["Predictors", "21"],
            ["Target", "Diabetes_binary"],
            ["Positive prevalence", "13.9%"],
            ["Imbalance ratio", "6.18:1"],
            ["Missing values", "0"],
            ["Exact duplicate rows", "24,206 (9.5%), retained"],
            ["Diagnostic laboratory variables", "None"],
        ],
        [2700, 1800],
        font_size=8.2,
    )
    subheading("B. Predictors and signal")
    body(
        "The variables span cardiometabolic status (high blood pressure, high cholesterol, body-mass index, stroke, and heart disease), behavior (smoking, physical activity, fruit and vegetable intake, heavy alcohol use), functional health (general health, mental and physical health days, difficulty walking), access and socioeconomic conditions (healthcare access, cost barriers, education, and income), and demographics (sex and age). The strongest univariate associations are general health, hypertension, difficulty walking, BMI, high cholesterol, and age, but no single survey item provides diagnostic-level separation."
    )
    subheading("C. Bias and measurement constraints")
    body(
        "BRFSS is cross-sectional and self-reported. It is vulnerable to recall, social-desirability, nonresponse, and telephone-coverage bias. The outcome represents previously diagnosed diabetes rather than biochemical confirmation, so undiagnosed cases may be labeled negative. Variables are coarsely quantized, and the unweighted public extract may not reproduce nationally representative estimates. These characteristics bound attainable performance and require cautious interpretation."
    )

    heading("Methodology")
    subheading("A. Protocol and leakage control")
    body(
        "All models followed one protocol. Data were divided using a stratified 60% training, 20% validation, and 20% test split, preserving prevalence and reserving 50,736 respondents for final evaluation. Numerical and ordinal features were standardized where required, binary features passed through unchanged, and all preprocessing transformers were fitted only within training folds. The test set was not used for hyperparameter selection, calibration fitting, or threshold design."
    )
    subheading("B. Models and tuning")
    body(
        "Nine estimators covered linear, instance-based, kernel, tree, boosting, and neural families: logistic regression [25], random forest [26], extra trees, histogram gradient boosting [27], XGBoost [28], LightGBM [29], multilayer perceptron, k-nearest neighbors [30], and support-vector machine [31]. Randomized hyperparameter search used three-fold stratified cross-validation on the training partition with ROC-AUC as the objective. KNN and SVM were tuned on stratified subsamples for computational tractability and then evaluated on the complete test set."
    )
    subheading("C. Metrics and statistical comparison")
    body(
        "Models were ranked by held-out ROC-AUC. We also report PR-AUC, MCC, balanced accuracy, F1, precision, sensitivity, specificity, accuracy, and Brier score. For each model, a 95% confidence interval for test ROC-AUC was estimated with 1,000 nonparametric bootstrap resamples [32]. Pairwise AUC differences used the same method in a paired implementation written in Python: each replicate sampled respondent indices once with replacement and applied the same indices to both prediction vectors, preserving pairing. Practical relevance was judged by effect magnitude rather than statistical significance alone."
    )
    subheading("D. Calibration, risk bands, and explanation")
    body(
        "XGBoost was selected before test evaluation because it achieved the highest validation ROC-AUC (0.8283). It was assessed in its original form and after Platt and isotonic calibration. Brier score and expected calibration error (ECE, ten bins) quantified probabilistic quality. Four candidate thresholding strategies were compared on the validation partition, including population percentiles, quartiles, sensitivity-oriented ROC criteria, and validation-set operating points. The final 0.09/0.20/0.40 cut points were then evaluated once on the untouched test set. SHAP and permutation importance characterized global signal without assigning causal meaning."
    )

    heading("Results")
    subheading("A. Model comparison")
    body(
        "Histogram gradient boosting and XGBoost tied for the highest test discrimination (ROC-AUC 0.8268), followed closely by LightGBM (0.8265), random forest (0.8237), and MLP (0.8232). Regularized logistic regression reached 0.8197, only 0.007 below the leaders. KNN achieved 0.7947 and SVM 0.6834. The best PR-AUC was 0.4237, about three times the random baseline equal to prevalence (0.139), showing useful but still modest absolute signal."
    )
    add_figure(doc, auc_chart, 3.05, "Figure 1. Held-out ROC-AUC for the nine estimators. Values come from the untouched test partition.", roles)
    caption("Table II. Held-out test performance ranked by ROC-AUC; CI values use 1,000 bootstrap resamples.")
    add_scientific_table(
        doc,
        ["Model", "ROC-AUC", "95% CI", "PR-AUC", "MCC", "Brier"],
        rows,
        [1050, 600, 1000, 600, 600, 650],
        font_size=6.8,
    )
    body(
        "The paired difference between the best model and logistic regression was ΔAUC = 0.007 (bootstrap 95% CI: 0.006-0.009). The interval excludes zero because the test set is large, but a seven-thousandth AUC gain is unlikely to change clinical decisions by itself. The leading boosting models were mutually indistinguishable. This supports H1 in its clinically relevant form: there is a statistically detectable difference, but it is too small to justify complexity on discrimination alone."
    )
    body(
        "Raw accuracy was especially misleading. Tree-based models reported approximately 0.866 accuracy, barely above the majority-class baseline of 0.861, while detecting only 4%-16% of positive cases at a 0.5 cutoff. Class-weighted logistic regression achieved the highest MCC (0.356), balanced accuracy (0.744), F1 (0.441), and sensitivity (0.760). Threshold choice and calibration therefore matter more than the default classifier label."
    )
    subheading("B. Probability calibration")
    body(
        "The original boosted model was already well calibrated (test ECE 0.0044; Brier 0.0974). Isotonic calibration reduced ECE to 0.0029 while retaining a comparable Brier score of 0.0975. Platt scaling degraded calibration in this large sample (ECE 0.0285). Isotonic probabilities were therefore used for risk stratification."
    )
    caption("Table III. Calibration of the selected XGBoost model. Lower values are better.")
    add_scientific_table(
        doc,
        ["Variant", "Brier val.", "ECE val.", "Brier test", "ECE test"],
        [
            ["Original", "0.0968", "0.0040", "0.0974", "0.0044"],
            ["Platt", "0.0985", "0.0257", "0.0995", "0.0285"],
            ["Isotonic", "0.0965", "0.0000", "0.0975", "0.0029"],
        ],
        [900, 900, 900, 900, 900],
        font_size=7.2,
    )
    add_figure(doc, calibration_chart, 3.05, "Figure 2. Held-out Brier score and expected calibration error before and after calibration.", roles)
    subheading("C. Risk stratification and held-out assessment")
    body(
        "The validation-selected cut points produced monotonic separation on the held-out test set. Observed prevalence increased from 2.9% in the Low band to 14.2% in Moderate, 29.1% in High, and 48.0% in Very high. Relative to Low risk, the odds ratios were 5.5, 13.7, and 30.7 with narrow confidence intervals. This sixteen-fold prevalence gradient suggests potential utility for population triage while requiring prospective external validation before clinical use."
    )
    caption("Table IV. Observed outcomes by calibrated risk band in the held-out test set (n = 50,736).")
    add_scientific_table(
        doc,
        ["Band", "n", "Prev.", "OR (95% CI)", "RR"],
        [
            ["Low", "26,149", "2.9%", "1.0 (ref)", "1.0"],
            ["Moderate", "11,877", "14.2%", "5.5 (5.0-6.0)", "4.9"],
            ["High", "7,844", "29.1%", "13.7 (12.5-14.9)", "10.0"],
            ["Very high", "4,866", "48.0%", "30.7 (28.0-33.6)", "16.4"],
        ],
        [850, 650, 850, 1450, 700],
        font_size=7.0,
    )
    add_figure(doc, risk_or_chart, 3.05, "Figure 3. Odds ratios and 95% confidence intervals by calibrated risk band on a logarithmic scale.", roles)
    add_figure(doc, risk_distribution_chart, 3.05, "Figure 4. Population size and observed diabetes prevalence across the four risk bands.", roles)
    subheading("D. Interpretability")
    body(
        "SHAP attributions and permutation importance agreed on the dominant signal. General health, hypertension, BMI, age, and high cholesterol increased predicted risk, while higher income, physical activity, and fruit and vegetable intake were associated with lower risk. These patterns are clinically plausible and consistent with established cardiometabolic and socioeconomic risk, but they remain associations within the fitted model rather than causal effects."
    )
    add_figure(doc, assets["shap"], 3.05, "Figure 5. SHAP summary for the selected model. Color represents feature value and horizontal position represents contribution to predicted risk.", roles)
    add_figure(doc, permutation_chart, 3.05, "Figure 6. Permutation importance measured as the decrease in ROC-AUC. The ranking corroborates the SHAP signal.", roles)

    heading("Discussion")
    subheading("A. Why complex models provide little additional discrimination")
    body(
        "The five strongest estimators differ by less than 0.004 ROC-AUC, and the best nonlinear model exceeds logistic regression by only 0.007. This near-equivalence suggests that much of the BRFSS signal is captured by low-order structure. Once the main effects of general health, hypertension, BMI, age, and cholesterol are represented, the evaluated flexible models obtain little additional discrimination."
    )
    subheading("B. The observed performance plateau")
    body(
        "Seven of the nine estimators clustered between ROC-AUC 0.819 and 0.827, while KNN and SVM performed below that range. The clustering among the leading models is consistent with an information ceiling imposed by coarse self-reported predictors, not universal impossibility. Biomarkers, longitudinal trajectories, medication history, genetics, or more precise behavior measurements could shift that ceiling. The conclusion is comparative and limited to the present variables, models, and population."
    )
    subheading("C. Clinical and public-health meaning")
    body(
        "An AUC near 0.83 is insufficient for individual diagnosis but useful for population triage. The Very high band represents about 10% of the cohort and concentrates an observed prevalence of 48%, a 3.4-fold enrichment over the base rate. A questionnaire-based pre-screen could prioritize confirmatory HbA1c or glucose testing when resources are limited. Diagnosis still requires biochemical confirmation under clinical standards [33]."
    )
    body(
        "The actionable output is not a binary label but a calibrated probability, risk band, uncertainty statement, and explanation. Presenting the 0.5 default prediction would omit most cases for several models and encourage automation bias. Safe deployment requires local recalibration, monitoring for population drift, documentation of the intended screening use, and interfaces that clearly state that the result is not a diagnosis."
    )
    subheading("D. Leakage and data governance")
    body(
        "The feature set avoids direct diagnostic-label leakage from laboratory diagnostic criteria. Adding HbA1c or fasting glucose would change the task from pre-laboratory screening toward a diagnostic-adjacent setting and could improve discrimination, but the magnitude of that improvement was not estimated in this study. Variable admissibility is therefore at least as consequential as choosing among the leading estimators [8], [9], [33]."
    )
    subheading("E. Interpretability and reporting")
    body(
        "Agreement between SHAP and permutation importance increases confidence that the model uses recognizable epidemiological structure. It does not prove causality or fairness. Income and healthcare access may encode structural inequities, and explanations can appear plausible even when validation is inadequate [19]. Reporting should follow transparent guidance such as TRIPOD+AI, including data provenance, missingness, sample flow, calibration, thresholds, and external validation [35]."
    )

    heading("Lessons for Clinical Machine Learning")
    body(
        "First, strong baselines are mandatory: complex models should be compared against regularized logistic regression under the same preprocessing and splits. Second, leakage control begins with a clinical definition of what information is available at prediction time. Third, probability calibration is not optional when outputs drive thresholds. Fourth, interpretability must connect to decisions rather than decorative plots. Fifth, a model that supports screening must be governed, recalibrated, monitored, and prevented from replacing professional judgment."
    )

    heading("Limitations and Threats to Validity")
    body(
        "The study is observational and cross-sectional. BRFSS supports association, not causation, and SHAP values or odds ratios must not be interpreted as intervention effects. Predictors and outcome are self-reported by telephone, which introduces recall and desirability bias. Previously undiagnosed diabetes is labeled negative, attenuating measured associations and potentially underestimating attainable signal."
    )
    body(
        "Selection and coverage bias may underrepresent people without stable telephone access and nonrespondents. The public extract does not apply survey weights, so descriptive estimates should not be interpreted as nationally representative. Socioeconomic variables may reflect access disparities rather than biology, creating equity concerns if the score is used to ration care."
    )
    body(
        "The dataset lacks HbA1c, glucose, lipids, detailed family history, medication, genetics, and longitudinal measurements. The observed plateau is conditional on this restricted feature space. Results also concern United States adults surveyed in 2015; transport to other countries, health systems, or current populations requires external validation and recalibration."
    )
    body(
        "Exact duplicate response vectors can cross random partitions and may inflate absolute discrimination. They were retained because respondent identifiers are unavailable and removing repeated categorical profiles would change prevalence. The magnitude and model-specific effect of this contamination are unknown; a grouped sensitivity analysis should be added when identifiers or defensible grouping keys become available."
    )
    body(
        "SVM and KNN were tuned on subsamples, which may underestimate their performance. Threshold-dependent metrics use a 0.5 cutoff and should not be confused with the calibrated risk-band policy. Finally, the study uses one stratified split with bootstrap uncertainty. Repeated cross-validation, multiple seeds, prospective validation, and decision-curve analysis remain important extensions."
    )

    heading("Future Work")
    body(
        "Future work should validate the model prospectively in Mexican and local clinical populations, monitor calibration over time, and evaluate fairness across sex, age, income, and access groups. Longitudinal PREDIA records could support trajectory models that distinguish stable risk from rapid deterioration. Patient-level explanations should be tested with clinicians for comprehension and actionability. Model and dataset versioning, automated metric gates, drift monitoring, and audit logs are needed before operational use."
    )

    heading("Conclusions")
    body(
        "Across nine estimators, available information limited performance more than algorithm choice. Seven models clustered between ROC-AUC 0.819 and 0.827, and the best nonlinear approach improved over regularized logistic regression by only ΔAUC = 0.007. This difference was statistically detectable but practically small. Raw accuracy obscured poor sensitivity at the default threshold, while calibration and validation-selected risk bands produced a clinically clearer output."
    )
    body(
        "Isotonic calibration achieved test ECE = 0.0029, and the four risk bands separated observed prevalence from 2.9% to 48.0%. SHAP and permutation importance recovered plausible cardiometabolic and socioeconomic patterns. The model is therefore best understood as a transparent population-screening aid that prioritizes confirmatory evaluation, not as a diagnostic system. Methodological rigor, admissible variables, calibration, interpretability, transparent reporting, and local monitoring offer more value than marginal gains in discrimination."
    )

    heading("Declarations")
    subheading("Conflict of interest")
    body("The authors declare that they have no known competing financial interests or personal relationships that could have influenced the work reported in this article.")
    subheading("Author contributions")
    body("Guillermo Álvarez Sánchez: conceptualization, methodology, software, data curation, formal analysis, model implementation, visualization, original draft, and review and editing.")
    body("Iván Peredo Valderrama: scientific supervision, methodological guidance, validation, correspondence, and review and editing.")
    body("José Javier Moya Moya: academic and scientific supervision, methodological review, validation, and review and editing.")
    body("Jairo Pacindo Piña: validation of results, critical review of scientific content, and review and editing.")
    subheading("Data and materials availability")
    body("The BRFSS-derived data used in this study are publicly available [23], [24]. Additional derived materials are available from the corresponding author upon reasonable request.")
    subheading("Funding")
    body("This research received no specific external funding from a public, commercial, or not-for-profit funding agency.")
    subheading("Acknowledgements")
    body("The authors thank the Universidad Politécnica de Querétaro for the facilities and institutional support provided during this research.")

    heading("References")
    for reference in parse_references():
        add_role_paragraph(
            doc,
            reference,
            roles["reference"],
            size=9.2,
            after=1.2,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )


def build() -> Path:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="cimerti-build-") as tmp:
        work_dir = Path(tmp)
        working = work_dir / "working.docx"
        shutil.copy2(REFERENCE, working)
        doc = Document(working)
        patch_first_page(doc)
        assets = extract_template_images(work_dir)
        roles = strip_scientific_body(doc)
        build_scientific_body(doc, roles, assets, work_dir)
        doc.core_properties.title = (
            "How Much Does Machine Learning Really Improve Diabetes Risk Prediction "
            "from Population Health Indicators?"
        )
        doc.core_properties.author = (
            "Guillermo Álvarez Sánchez; Iván Peredo Valderrama; "
            "José Javier Moya Moya; Jairo Pacindo Piña"
        )
        doc.core_properties.last_modified_by = "PREDIA / Universidad Politécnica de Querétaro"
        doc.save(OUTPUT)
        patch_footer_package(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
