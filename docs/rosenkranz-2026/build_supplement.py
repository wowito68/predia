"""Build the anonymized supplement from immutable analysis outputs."""
from __future__ import annotations

import csv
import hashlib
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_manuscript import HERE, RESULTS, add_rich_text, add_table, configure_document

OUTPUT = HERE / "suplemento-metodologico-ensanut.docx"
DESIGN_PATH = RESULTS / "design_sensitivity/analysis.json"
DATA = json.loads(DESIGN_PATH.read_text())
POLICY = json.loads((RESULTS / "postlock_policy_analysis.json").read_text())
PRIMARY = json.loads((RESULTS / "evaluation_2021.json").read_text())
SELECTED = next(m for m in PRIMARY["experiments"]["core"]["models"]
                if m["family"] == PRIMARY["experiments"]["core"]["recommended_family"])
DISPLAY = {(r["age"], r["metric"]): r for r in csv.DictReader(
    (HERE / "figures/age-omissions-logit-ci.csv").open())}
YEARS = (2012, 2016, 2018, 2021)


def prose(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_rich_text(p, text)
    return p


def page(doc, title):
    doc.add_page_break()
    doc.add_paragraph(title, "Heading 1")


def bounds(row, digits=3, scale=1):
    return f"{float(row['lower'])*scale:.{digits}f} a {float(row['upper'])*scale:.{digits}f}"


def estimate_ci(row, digits=1, scale=1):
    return f"{float(row['estimate'])*scale:.{digits}f}\n({bounds(row, digits, scale)})"


def missing_rows(variables):
    lookup = {(r["wave"], r["predictor"]): r for r in DATA["missingness"]}
    rows = []
    for key, name in variables:
        values = [name]
        for year in YEARS:
            r = lookup[year, key]
            if year == 2012 and key in ("systolic_bp", "diastolic_bp", "hba1c_pct"):
                values.append("No armonizado")
            else:
                values.append(f"{r['missing']:,} ({r['missing_pct']:.1f}%)\nP: {r['missing_weighted_pct']:.1f}%")
        rows.append(values)
    return rows


def build():
    doc = Document()
    configure_document(doc)
    title = "Suplemento metodológico de la priorización de pruebas con ENSANUT"
    doc.core_properties.title = title
    doc.core_properties.subject = "Análisis complementario de encuesta y omisiones"
    doc.add_paragraph(title, "Title")
    prose(doc, "Este suplemento acompaña la validación temporal del modelo de seis variables "
          "no invasivas para priorizar pruebas de disglucemia en adultos mexicanos. Detalla "
          "la construcción de las cohortes, los datos faltantes y la sensibilidad de la "
          "inferencia al diseño de encuesta. El contraste amplía la incertidumbre: no "
          "convierte una ventaja de discriminación en superioridad operativa demostrada.")
    doc.add_paragraph("Alcance temporal", "Heading 1")
    prose(doc, "El desarrollo utilizó 2012 y 2016; la selección de configuración y umbral, "
          "2018; y la evaluación bloqueada, 2021. La factibilidad y prevalencia agregada de "
          "2021 habían sido inspeccionadas; no se afirma prerregistro externo ni "
          "desconocimiento absoluto de esa ola. El complemento de capacidad y este "
          "contraste de varianza se especificaron después de conocer el resultado "
          "primario, antes de sus respectivas ejecuciones. No se modificaron modelos, "
          "probabilidades, imputadores, umbral ni resultados anteriores.")
    doc.add_paragraph("Población y unidad de inferencia", "Heading 1")
    prose(doc, "El archivo bioquímico de 2021 contiene 2,034 registros con ponderador "
          "ponde_g positivo, 85 estratos y 397 unidades primarias de muestreo (UPM). Los "
          "1,724 adultos elegibles constituyen un dominio: sus ponderadores, estratos y "
          "UPM coinciden exactamente con el análisis bloqueado. Se conservan los demás "
          "registros con contribución linealizada cero, incluidas siete UPM sin adultos "
          "elegibles. Persisten 13 estratos con una sola UPM; conservar el dominio no "
          "recupera información inexistente del diseño original [1,3].")
    prose(doc, "El objetivo de sensibilidad ≥80% en 2018 fue una elección operativa para "
          "priorizar detección, no una norma clínica. Las capacidades de 200, 400 y 600 "
          "pruebas por 1,000 exploran otro balance entre cobertura y carga; no "
          "reemplazan el umbral bloqueado 0.238346 ni validan preferencias sanitarias.")

    page(doc, "Selección secuencial de participantes")
    step_names = [
        ("linked_rows", "Registros tras enlace"),
        ("age_20_to_110", "Edad de 20 a 110 años"),
        ("known_diabetes_status", "Estado de diabetes conocido"),
        ("without_prior_diabetes", "Sin diabetes diagnosticada"),
        ("not_currently_pregnant", "Sin embarazo registrado"),
        ("fasting_sample", "Submuestra de ayuno"),
        ("positive_survey_weight", "Ponderador positivo"),
        ("valid_fasting_glucose", "Glucosa válida y cohorte final"),
    ]
    lookup = {(r["wave"], r["step"]): r for r in DATA["cohort_flows"]}
    rows = []
    for key, label in step_names:
        row = [label]
        for year in YEARS:
            value = lookup[year, key]
            excluded = value["excluded_step"]
            row.append(f"{value['remaining']:,}" if excluded is None else
                       f"{value['remaining']:,}\n(-{int(excluded):,})")
        rows.append(row)
    add_table(doc, ["Etapa", *map(str, YEARS)], rows, [2.3, 1.1, 1.1, 1.1, 1.1],
              "Tabla S1. Flujo de las cuatro cohortes", "Cada celda muestra participantes "
              "restantes; entre paréntesis, exclusiones en ese paso. Los totales fuente "
              "de cuestionarios y laboratorios no son pasos secuenciales comparables.")
    prose(doc, "El enlace no comienza en el mismo universo en todas las olas: en 2018 "
          "parte del cuestionario de adultos y en 2012 de registros bioquímicos enlazados. "
          "La etapa de ayuno también excluye ausencia de información o de enlace "
          "bioquímico; sus exclusiones no equivalen a personas que declararon ayuno "
          "insuficiente. En 2012 no se verifican horas individuales. La ausencia de "
          "información de embarazo no motivó exclusión.")

    page(doc, "Datos faltantes del conjunto núcleo")
    prose(doc, "Los faltantes se cuantificaron antes de la imputación, dentro de cada "
          "cohorte final. El denominador es 8,710 en 2012; 3,308 en 2016; 11,093 en "
          "2018; y 1,724 en 2021. Los porcentajes ponderados utilizan el ponderador "
          "bioquímico de cada ola.")
    core = [("age", "Edad"), ("female", "Sexo registrado"),
            ("parent_diabetes", "Diabetes en padre o madre"),
            ("diagnosed_hypertension", "Hipertensión diagnosticada"),
            ("bmi", "Índice de masa corporal"), ("waist_cm", "Cintura")]
    add_table(doc, ["Variable", *map(str, YEARS)], missing_rows(core),
              [2.1, 1.15, 1.15, 1.15, 1.15], "Tabla S2. Faltantes de los seis predictores",
              "Primera línea: n faltante (% no ponderado). P: porcentaje ponderado. "
              "Los faltantes de distintas variables pueden coincidir en una persona; "
              "no deben sumarse para calcular participantes incompletos.")
    prose(doc, "El núcleo estuvo completo en 7,522, 2,835, 9,784 y 1,523 participantes, "
          "respectivamente. El análisis principal retuvo participantes con predictores "
          "incompletos mediante transformaciones aprendidas en desarrollo. En el modelo "
          "aditivo se emplearon medianas para continuas y moda con indicadores de "
          "ausencia para binarias. No se aprendieron imputadores con la ola de prueba.")

    page(doc, "Variables complementarias y límites de la imputación")
    add_table(doc, ["Variable", *map(str, YEARS)], missing_rows([
        ("systolic_bp", "Presión sistólica"), ("diastolic_bp", "Presión diastólica"),
        ("hba1c_pct", "HbA1c")]), [2.1, 1.15, 1.15, 1.15, 1.15],
        "Tabla S3. Faltantes fuera del conjunto núcleo",
        "Primera línea: n faltante (% no ponderado). P: porcentaje ponderado. "
        "No armonizado significa que el análisis no incorporó esa variable en 2012; "
        "no afirma que ninguna fuente de esa encuesta la hubiera medido.")
    prose(doc, "La presión arterial se utiliza en el modelo ampliado, entrenado solo "
          "con 2016. Por ello, su comparación con el núcleo también cambia las olas "
          "de entrenamiento y no estima aisladamente el efecto de añadir presión "
          "arterial. La HbA1c define desenlaces secundarios; no es un predictor. Su "
          "submuestra de 2012 no se consideró comparable.")
    prose(doc, "No se imputó glucosa ausente. Los porcentajes de las tablas S2 y S3 "
          "describen la información de participantes ya elegibles; no cuantifican "
          "por sí solos la no respuesta bioquímica entre todas las personas encuestadas. "
          "La tabla S1 hace visibles las pérdidas anteriores. La imputación simple "
          "no elimina un posible sesgo de selección ni incorpora incertidumbre "
          "por imputación; el análisis de casos completos tampoco lo descarta.")
    prose(doc, "La disglucemia primaria se definió como glucosa en ayuno ≥100 mg/dL. "
          "Una medición ≥126 mg/dL identifica rango diabético en la encuesta, "
          "no un diagnóstico confirmado. Estas estimaciones corresponden a la "
          "población elegible sin diagnóstico previo y no a incidencia futura.")

    page(doc, "Estimación de la varianza")
    prose(doc, "La implementación es propia en Python. Se usaron ponderadores "
          "bioquímicos y UPM anidadas en estratos, conservando el archivo completo "
          "con ponderador positivo. Para una razón, se divide el total ponderado "
          "del numerador entre el total ponderado del denominador. Su contribución "
          "individual es el ponderador por la diferencia entre el numerador individual "
          "y el producto de la razón estimada por el denominador individual, todo dividido "
          "entre el total ponderado del denominador [1].")
    prose(doc, "Para AUC se comparan todos los pares caso-control, ponderados por "
          "el producto de sus pesos, con medio crédito por empate. La contribución "
          "de un caso es su peso dividido entre el peso total de casos, multiplicado "
          "por su proporción ponderada de controles superados menos la AUC. Para "
          "un control se utiliza la proporción ponderada de casos que lo superan "
          "y el peso total de controles. La diferencia de AUC resta contribuciones "
          "en los mismos participantes, preservando su covarianza.")
    prose(doc, "Se suman contribuciones dentro de cada UPM. En un estrato con m UPM, "
          "su varianza es la suma de cuadrados de desviaciones de esos totales "
          "respecto a su media, multiplicada por m/(m−1). Se suman las contribuciones "
          "de estratos, sin corrección por población finita. Para estratos de una "
          "UPM se comparan tres supuestos [1,4]:")
    prose(doc, "**Contribución cero:** omite su varianza, sin afirmar que la UPM fue "
          "seleccionada con certeza. **Centrado global:** utiliza la desviación "
          "cuadrada frente a la media de totales de todas las UPM. **Promedio:** "
          "asigna a cada estrato único el promedio de las contribuciones de "
          "varianza de estratos con más de una UPM. El promedio se fijó como "
          "sensibilidad destacada antes de ejecutarla; se muestran los tres.")
    prose(doc, "Los IC95% usan el cuantil t con grados de libertad del dominio "
          "(UPM representadas menos estratos representados); en razones "
          "condicionadas, se usa el dominio del denominador. Son 305 para la "
          "AUC completa y 288 para su contraste común. El jackknife elimina "
          "una UPM por réplica y multiplica por m/(m−1) los pesos de las "
          "restantes de su estrato. Acumula (m−1)/m por la desviación cuadrada "
          "respecto a la estimación completa: 384 réplicas, con estratos únicos "
          "fijos [2,5]. Ningún procedimiento reestima el modelo ni reproduce "
          "todas las etapas y ajustes del diseño original.")

    page(doc, "Sensibilidad de los resultados principales")
    rows = []
    for method, label in (("omit", "Contribución cero"), ("center", "Centrado global"),
                          ("average", "Promedio")):
        rows.append([label, bounds(DATA["metrics"]["roc_auc"][method]),
                     bounds(DATA["metrics"]["common_auc_difference"][method])])
    rows.append(["Jackknife", bounds(DATA["jackknife"]["metrics"]["roc_auc"]),
                 bounds(DATA["jackknife"]["metrics"]["common_auc_difference"])])
    rows.insert(0, ["Bootstrap original", bounds(SELECTED["bootstrap_ci"]["roc_auc"]), "0.008 a 0.046"])
    add_table(doc, ["Procedimiento", "IC95% AUC", "IC95% diferencia"], rows,
              [2.4, 2.1, 2.2], "Tabla S4. Intervalos de discriminación",
              "AUC del modelo: 0.694 (n=1,724). Diferencia modelo menos puntaje: "
              "0.027 (n=1,413). Los puntos no cambian. La comparación de 2021 "
              "no es validación externa del puntaje desarrollado con 2021-2023.")
    metrics = [("Sensibilidad (%)", "sensitivity", 100, 1),
               ("Especificidad (%)", "specificity", 100, 1),
               ("Remisión (%)", "referral_fraction", 100, 1),
               ("Brier", "brier", 1, 3),
               ("Omisiones por 1,000", "missed_per_1000", 1, 1)]
    rows = [[label, f"{DATA['metrics'][key]['estimate']*scale:.{digits}f}",
             bounds(DATA["metrics"][key]["average"], digits, scale)]
            for label, key, scale, digits in metrics]
    add_table(doc, ["Métrica", "Estimación", "IC95% promedio"], rows,
              [2.5, 1.8, 2.4], "Tabla S5. Otras métricas con linealización",
              "IC t simétricos del análisis complementario con tratamiento promedio. "
              "El manuscrito conserva los intervalos originales en su tabla 2.")

    page(doc, "Omisiones por edad")
    rows = []
    for age in ("20-34", "35-54", "55+"):
        old = POLICY["estimates"][f"subgroup/age_group/{age}"]["missed_per_1000"]
        within = DISPLAY[age, "missed_within_1000"]
        whole = DISPLAY[age, "missed_total_1000"]
        rows.append([age, estimate_ci(old), estimate_ci(within), estimate_ci(whole)])
    add_table(doc, ["Edad", "Dentro del grupo\nBootstrap", "Dentro del grupo\nLogit-t", "Cohorte total\nLogit-t"],
              rows, [.7, 2.0, 2.0, 2.0], "Tabla S6. Alteraciones omitidas por 1,000 adultos",
              "Estimación e IC95%. Dentro del grupo: denominador de la misma edad. "
              "Cohorte total: contribución al total de adultos elegibles. "
              "Bootstrap: complemento previo de capacidad. Logit-t: linealización "
              "con tratamiento promedio de estratos únicos.")
    prose(doc, "Las contribuciones 28.4, 15.9 y 3.3 suman las 47.6 omisiones por "
          "1,000 de la cohorte completa. Los valores 73.0, 41.6 y 14.3 no deben "
          "sumarse porque tienen denominadores diferentes. El grupo de 20-34 "
          "años reúne 59.6% de las omisiones ponderadas: IC95% bootstrap previo "
          "47.1-72.4%; linealización promedio simétrica, 42.3-77.0%.")
    prose(doc, "Los intervalos simétricos del complemento se conservaron íntegros. "
          "En ≥55 años, el intervalo dentro del grupo fue −8.9 a 37.5 por 1,000 "
          "y el de contribución total, −2.0 a 8.6. Sus límites negativos reflejan "
          "la aproximación de Wald, no omisiones negativas. Tras inspeccionarlos "
          "se eligió la transformación logit-t para la figura y esta tabla: es "
          "una decisión de presentación posterior, no preespecificada.")
    prose(doc, "Se expresa la tasa como proporción y se divide su error estándar "
          "entre el producto de la proporción por su complemento; se aplica "
          "el cuantil t en escala logit y se transforma de regreso, multiplicando "
          "por 1,000. No se recortaron límites a cero ni cambiaron las estimaciones "
          "o varianzas. Los intervalos asimétricos amplios refuerzan la cautela "
          "en grupos pequeños; no demuestran causalidad ni justifican nuevos "
          "umbrales por edad.")

    page(doc, "Incertidumbre de la curva de decisión")
    selected_curves = [r for r in DATA["decision_curve"] if any(
        abs(r["threshold"]-cutoff) < 1e-12 for cutoff in (.1, .2, DATA["threshold"], .3))]
    rows = []
    for r in selected_curves:
        name = f"{r['threshold']:.2f}" if r["threshold"] != DATA["threshold"] else "0.238346"
        def dca_ci(key):
            return estimate_ci({"estimate": r[key]["estimate"], **r[key]["average"]}, 4)
        rows.append([name, dca_ci("model"), dca_ci("difference_all")])
    add_table(doc, ["Umbral", "Modelo menos nadie", "Modelo menos todos"], rows,
              [1.1, 2.8, 2.8], "Tabla S7. Beneficio neto y contrastes emparejados",
              "Estimación e IC95% puntual. Linealización con tratamiento promedio "
              "de estratos únicos. El beneficio de remitir a nadie es cero; "
              "la primera columna de resultados coincide con el beneficio del modelo.")
    prose(doc, "El beneficio neto divide verdaderos positivos ponderados entre "
          "población ponderada y resta la fracción de falsos positivos multiplicada "
          "por la razón entre umbral y su complemento [6]. Se calculó en 0.05-0.50, "
          "con incrementos de 0.01, más el umbral bloqueado. Los contrastes "
          "usan las diferencias individuales en la misma muestra; no restan "
          "intervalos independientes.")
    prose(doc, "En el umbral original, el IC del beneficio del modelo incluye "
          "cero bajo el supuesto promedio. Con contribución cero fue "
          "0.0026-0.0661 y con centrado global 0.0013-0.0673: la inferencia "
          "frente a no remitir es sensible a esos supuestos. La diferencia "
          "frente a remitir a todos conserva un intervalo positivo. No hay "
          "evidencia para afirmar superioridad frente a ambas alternativas "
          "en todo el intervalo de umbrales.")
    prose(doc, "Los IC son puntuales, no bandas simultáneas ni ajustes por "
          "multiplicidad. La curva no constituye beneficio clínico observado, "
          "evaluación económica ni autorización para usar probabilidades "
          "sobreestimadas como riesgo individual calibrado. Cambiar calibración "
          "o umbral requiere datos de actualización y evaluación independientes.")

    page(doc, "Trazabilidad y referencias metodológicas")
    prose(doc, "Se conservaron las huellas del estudio, la evaluación primaria "
          "y el modelo. El complemento es condicional al modelo fijo; no "
          "cuantifica incertidumbre del entrenamiento, todas las etapas de "
          "muestreo, ajustes de no respuesta ni calibración de ponderadores. "
          "Las comprobaciones numéricas no sustituyen revisión estadística "
          "independiente ni validación clínica prospectiva.")
    prose(doc, "La suite científica completó 24 pruebas, incluidas derivadas "
          "por diferencias finitas, AUC ponderada con empates, invariancia "
          "de escala de ponderadores, varianza de muestreo aleatorio simple, "
          "dominios con UPM de contribución cero y tratamientos de estratos "
          "únicos. La aproximación se implementó en Python, no mediante el "
          "paquete survey de R; su documentación se consultó como referencia.")
    prose(doc, "Los archivos de resultados conservan hashes del addendum, "
          "script y evaluación primaria. Las tablas de exclusión, faltantes "
          "e intervalos proceden de esas salidas; no se reestimó el estudio "
          "para componer este documento.")
    doc.add_paragraph("Referencias", "Heading 1")
    references = [
        "1. Lohr SL. Sampling: Design and Analysis. 3rd ed. Chapman & Hall/CRC; 2022. ISBN 9780367279509.",
        "2. Efron B, Tibshirani RJ. An Introduction to the Bootstrap. Chapman and Hall/CRC; 1993. doi:10.1201/9780429246593.",
        "3. CDC. NHANES Tutorials: Variance Estimation. https://wwwn.cdc.gov/nchs/nhanes/tutorials/varianceestimation.aspx. Consulta: 12 de septiembre de 2026. Referencia general de dominios; no sustituye el diseño ENSANUT.",
        "4. Lumley T. survey options. Documentación de survey. https://r-survey.r-forge.r-project.org/survey/html/surveyoptions.html. Consulta: 12 de septiembre de 2026.",
        "5. Lumley T. as.svrepdesign. Documentación de survey. https://r-survey.r-forge.r-project.org/survey/html/as.svrepdesign.html. Consulta: 12 de septiembre de 2026.",
        "6. Vickers AJ, Elkin EB. Decision curve analysis: a novel method for evaluating prediction models. Med Decis Making. 2006;26(6):565-574. doi:10.1177/0272989X06295361.",
    ]
    for reference in references:
        prose(doc, reference).alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "tables": len(doc.tables),
                      "input_sha256": hashlib.sha256(DESIGN_PATH.read_bytes()).hexdigest()}, indent=2))


if __name__ == "__main__":
    build()
