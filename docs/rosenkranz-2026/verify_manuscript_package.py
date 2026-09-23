"""Check document tables, source integrity, anonymity and final PDF structure."""
import hashlib
import json
import math
import re
from datetime import datetime, timezone

import pdfplumber
from docx import Document

from build_manuscript import HERE, RESULTS, SOURCE, source_word_counts
from build_supplement import DATA, POLICY, PRIMARY, SELECTED, missing_rows


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(table):
    return [[c.text for c in r.cells] for r in table.rows[1:]]


def main():
    project = RESULTS.parent
    for path, expected in (
        (RESULTS / "evaluation_2021.json", "4ba5a77643e1817692a0c0016d118007255fe1af7a3eb8995a0ba0cb84d75f9a"),
        (RESULTS / "models/core_spline_logistic.joblib", "c0d26e3175f76f539bff0e669c484f600e1e17c815d79e7f3ba8d43254bd6fb3"),
        (project / "postlock_design_sensitivity.py", DATA["script_sha256"]),
        (project / "DESIGN_SENSITIVITY_ADDENDUM.md", DATA["addendum_sha256"]),
    ):
        assert sha(path) == expected, path
    assert math.isclose(DATA["metrics"]["roc_auc"]["estimate"], SELECTED["test"]["roc_auc"], abs_tol=1e-12)
    assert math.isclose(sum(DATA["metrics"][f"age/{g}/missed_total_1000"]["estimate"]
                            for g in ("20-34", "35-54", "55+")),
                        DATA["metrics"]["missed_per_1000"]["estimate"], abs_tol=1e-12)
    for year in (2012, 2016, 2018, 2021):
        flow = [r for r in DATA["cohort_flows"] if r["wave"] == year]
        assert flow[-1]["remaining"] == POLICY["cohorts"][str(year)]["n"]
        for previous, current in zip(flow, flow[1:]):
            assert previous["remaining"]-current["remaining"] == current["excluded_step"]

    manuscript = Document(HERE / "manuscrito-ensanut-fortalecido.docx")
    supplement = Document(HERE / "suplemento-metodologico-ensanut.docx")
    assert len(manuscript.tables) == 3 and len(manuscript.inline_shapes) == 3
    assert len(supplement.tables) == 7
    metrics = [("roc_auc", 3, 1), ("pr_auc", 3, 1), ("brier", 3, 1),
               ("calibration_intercept", 3, 1), ("calibration_slope", 3, 1),
               ("sensitivity", 1, 100), ("specificity", 1, 100), ("ppv", 1, 100),
               ("npv", 1, 100), ("referral_fraction", 1, 100),
               ("cases_detected_per_1000", 1, 1), ("tests_per_case_detected", 2, 1)]
    for row, (key, digits, scale) in zip(rows(manuscript.tables[1]), metrics):
        fmt = lambda v: f"{v*scale:.{digits}f}" + ("%" if scale == 100 else "")
        ci = SELECTED["bootstrap_ci"][key]
        assert row[1:] == [fmt(SELECTED["test"][key]), f"{fmt(ci['lower'])} a {fmt(ci['upper'])}"]
    for row, capacity in zip(rows(manuscript.tables[2]), (.2, .4, .6)):
        assert row[0] == str(round(capacity*1000))
        for text, key in zip(row[1:], (f"capacity/{capacity}/model", f"capacity/{capacity}/published", f"difference/{capacity}")):
            r = POLICY["estimates"][key]["detected_per_1000"]
            assert text == f"{r['estimate']:.1f}\n({r['lower']:.1f} a {r['upper']:.1f})"
    for table_index, variables in ((1, [("age", "Edad"), ("female", "Sexo registrado"),
                                        ("parent_diabetes", "Diabetes en padre o madre"),
                                        ("diagnosed_hypertension", "Hipertensión diagnosticada"),
                                        ("bmi", "Índice de masa corporal"), ("waist_cm", "Cintura")]),
                                   (2, [("systolic_bp", "Presión sistólica"),
                                        ("diastolic_bp", "Presión diastólica"), ("hba1c_pct", "HbA1c")])):
        assert rows(supplement.tables[table_index]) == missing_rows(variables)
    common = PRIMARY["paired_auc_comparisons"]["core_minus_published_score_auc_common_sample"]
    assert rows(supplement.tables[3])[0][2] == f"{common['lower']:.3f} a {common['upper']:.3f}"

    report = {"checked_at_utc": datetime.now(timezone.utc).isoformat(),
              "counts": source_word_counts(SOURCE.read_text()), "files": {},
              "scientific_table_checks": "passed", "frozen_inputs": "unchanged"}
    assert report["counts"]["abstract"] <= 250 and report["counts"]["all_before_references"] <= 3500
    for stem, document, page_count in (("manuscrito-ensanut-fortalecido", manuscript, 17),
                                        ("suplemento-metodologico-ensanut", supplement, 9)):
        assert document.core_properties.author == "Anonymous"
        assert document.core_properties.last_modified_by == "Anonymous"
        assert not document.element.xpath(".//w:ins|.//w:del|.//w:commentRangeStart")
        assert not document.styles["Title"]._element.xpath(".//w:pBdr")
        paragraphs = list(document.paragraphs) + [p for table in document.tables for row in table.rows for c in row.cells for p in c.paragraphs]
        assert all((r.font.size or p.style.font.size or document.styles["Normal"].font.size).pt == 11
                   for p in paragraphs for r in p.runs if r.text.strip())
        pdf = HERE / f"{stem}.pdf"
        with pdfplumber.open(pdf) as rendered:
            assert len(rendered.pages) == page_count
            text = "\n".join(p.extract_text() or "" for p in rendered.pages)
            assert not re.search(r"Guillermo|Arianna|Cristopher|Gabriel|/home/wowo|codex-file|turn\d+search|\ufffd", text, re.I)
            assert "Priorización" in text or "priorización" in text
            assert all(p.width == 612 and p.height == 792 for p in rendered.pages)
        report["files"][stem] = {"docx_sha256": sha(HERE / f"{stem}.docx"),
                                  "pdf_sha256": sha(pdf), "pages": page_count}
    (HERE / "package-verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
