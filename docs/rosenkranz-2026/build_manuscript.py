from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = HERE / "manuscrito-ciego.md"
OUTPUT = HERE / "manuscrito-ensanut-fortalecido.docx"
RESULTS = ROOT / "ml-research" / "ensanut_temporal" / "results"
QA_SCREEN = HERE / "evidence/mobile-dark/03-tamizaje-resultado.png"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=11, bold=None, italic=None, color="000000") -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_rich_text(paragraph, text: str, size=11) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text.replace(">=", "≥"))
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part.replace("  ", " "))
            set_run_font(run, size=size)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, size=11, color="595959")


def configure_document(document: Document) -> None:
    for style in document.styles:
        for border in style._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(11)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.5

    heading_sizes = {1: 11, 2: 11, 3: 11}
    for level, size in heading_sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if level < 3 else 7)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.5

    if "Figure Caption" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = document.styles["Figure Caption"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(5)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = False

    footer = section.footer.paragraphs[0]
    add_page_field(footer)

    document.core_properties.author = "Anonymous"
    document.core_properties.last_modified_by = "Anonymous"
    document.core_properties.title = "Tamizaje no invasivo de disglucemia no diagnosticada en adultos mexicanos"
    document.core_properties.subject = "Manuscrito ciego de investigación epidemiológica"
    document.core_properties.keywords = "ENSANUT, tamizaje, diabetes, validación temporal"


def add_table(document: Document, headers, rows, widths, caption: str, note: str) -> None:
    caption_p = document.add_paragraph(style="Figure Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption_p.add_run(caption)
    set_run_font(run, size=11, bold=True)
    caption_p.paragraph_format.keep_with_next = True

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, width)
        set_cell_shading(cell, "333333")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(header)
        set_run_font(run, size=11, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[index]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F2F2F2")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.5
            run = paragraph.add_run(str(value))
            set_run_font(run, size=11)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(3)
    set_run_font(paragraph.add_run(note), size=11, italic=True)


def add_figure(document: Document, path: Path, width: float, caption: str, page_break=True) -> None:
    if page_break:
        document.add_page_break()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_p = document.add_paragraph(style="Figure Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich_text(caption_p, caption, size=11)


def add_marker(document: Document, marker: str, mobile: Path) -> None:
    if marker == "[TABLE1]":
        document.add_page_break()
        policy = json.loads((RESULTS / "postlock_policy_analysis.json").read_text())
        cohort_rows = [[year, "Desarrollo" if year in ("2012", "2016") else "Selección" if year == "2018" else "Prueba bloqueada", f"{row['n']:,}", f"{row['complete_core']:,}", f"{100*row['weighted_prevalence']:.1f}%"] for year, row in policy["cohorts"].items()]
        add_table(
            document,
            ["Ola", "Función", "Elegibles", "Núcleo completo", "Prevalencia ponderada"],
            cohort_rows,
            [0.65, 1.50, 1.05, 1.55, 1.95],
            "Tabla 1. Cohortes analíticas por ola ENSANUT",
            "Nota. Datos completos: seis predictores núcleo disponibles. El análisis principal conservó participantes mediante imputación aprendida en desarrollo. Prevalencia: glucosa ≥100 mg/dL, sin diagnóstico previo.",
        )
    elif marker == "[TABLE2]":
        document.add_page_break()
        evaluation = json.loads((RESULTS / "evaluation_2021.json").read_text())
        experiment = evaluation["experiments"]["core"]
        selected = next(item for item in experiment["models"] if item["family"] == experiment["recommended_family"])
        metrics = [
            ("ROC-AUC", "roc_auc", 3, False),
            ("Precisión promedio (AP)", "pr_auc", 3, False),
            ("Brier", "brier", 3, False),
            ("Intercepto de calibración", "calibration_intercept", 3, False),
            ("Pendiente de calibración", "calibration_slope", 3, False),
            ("Sensibilidad", "sensitivity", 1, True),
            ("Especificidad", "specificity", 1, True),
            ("Valor predictivo positivo", "ppv", 1, True),
            ("Valor predictivo negativo", "npv", 1, True),
            ("Fracción remitida", "referral_fraction", 1, True),
            ("Casos detectados por 1,000", "cases_detected_per_1000", 1, False),
            ("Pruebas por caso", "tests_per_case_detected", 2, False),
        ]
        rows = []
        for label, key, decimals, percent in metrics:
            def formatted(value):
                return f"{value * (100 if percent else 1):.{decimals}f}" + ("%" if percent else "")
            ci = selected["bootstrap_ci"][key]
            rows.append([label, formatted(selected["test"][key]), f"{formatted(ci['lower'])} a {formatted(ci['upper'])}"])
        add_table(
            document,
            ["Métrica", "Estimación", "IC95%"],
            rows,
            [2.85, 1.4, 2.45],
            "Tabla 2. Desempeño del modelo bloqueado en ENSANUT 2021",
            "Nota. n=1,724; 397 eventos. IC95% aproximados mediante 1,000 remuestreos de UPM dentro de estratos. Umbral bloqueado: 0.238346. Las pruebas por caso representan una estimación de remisión, no intervenciones realizadas.",
        )
    elif marker == "[TABLE3]":
        document.add_page_break()
        policy = json.loads((RESULTS / "postlock_policy_analysis.json").read_text())["estimates"]
        rows = []
        fmt = lambda metric: f"{metric['estimate']:.1f}\n({metric['lower']:.1f} a {metric['upper']:.1f})"
        for capacity in (0.2, 0.4, 0.6):
            rows.append([str(round(capacity*1000)), fmt(policy[f"capacity/{capacity}/model"]["detected_per_1000"]), fmt(policy[f"capacity/{capacity}/published"]["detected_per_1000"]), fmt(policy[f"difference/{capacity}"]["detected_per_1000"])])
        add_table(document, ["Pruebas por 1,000", "Detección con modelo", "Detección con puntaje", "Diferencia modelo - puntaje"], rows, [1.0, 1.9, 1.9, 1.9], "Tabla 3. Detección estimada a igual capacidad de laboratorio", "Nota. Análisis exploratorio posterior al bloqueo. n=1,413; 338 eventos; prevalencia ponderada 22.2%. Valores por 1,000 adultos; IC95% aproximados entre paréntesis. Desempate aleatorio esperado en el nivel límite. El comparador fue desarrollado con ENSANUT 2021-2023, por lo que esta evaluación no constituye su validación externa independiente. Ninguno de los tres IC de la diferencia excluye cero.")
    elif marker == "[FIGURE1]":
        add_figure(document, HERE / "figures/calibration-decision.png", 6.5,
                   "**Figura 1. Calibración y utilidad de decisión en ENSANUT 2021.** A: estimaciones puntuales de calibración. B: beneficio neto; cero representa remitir a nadie. C: contrastes emparejados, con detalle entre 0.05 y 0.30. B y C: IC95% puntuales de linealización con tratamiento promedio de estratos únicos, no bandas simultáneas. El modelo y el umbral original permanecen fijos.")
    elif marker == "[FIGURE2]":
        add_figure(document, RESULTS / "figures" / "capacity_comparison_2021.png", 6.7, "**Figura 2. Detección a igual capacidad de laboratorio.** Análisis exploratorio sobre la muestra común (n=1,413). IC95% aproximados por remuestreo de UPM dentro de estratos. La selección aleatoria representa su valor esperado. Las posiciones horizontales se separan ligeramente para facilitar la lectura. Los contrastes emparejados se presentan en la tabla 3.")
    elif marker == "[FIGURE3]":
        add_figure(
            document,
            HERE / "figures/age-omissions.png",
            6.7,
            "**Figura 3. Omisiones por edad con dos denominadores.** Izquierda: por 1,000 adultos del mismo grupo. Derecha: contribución por 1,000 adultos de la cohorte completa; las tres estimaciones suman 47.6. IC95% aproximados logit-t a partir de linealización con tratamiento promedio de estratos únicos. Las escalas horizontales difieren. Son estimaciones de remisión, no resultados de una intervención.",
        )
    elif marker == "[FIGURE4]":
        add_figure(
            document,
            mobile,
            2.8,
            "**Figura 4. Traducción del modelo a un prototipo móvil.** Captura con nombre y datos de un paciente simulado. La interfaz comunica prioridad de confirmación y el límite de no diagnóstico; el resultado no se persiste en el expediente.",
        )


def source_word_counts(text: str) -> dict[str, int]:
    abstract_match = re.search(r"## Resumen\s+(.*?)\s+\*\*Palabras clave:", text, re.S)
    main_match = re.search(r"## Introducción\s+(.*?)\s+## Referencias", text, re.S)
    clean = lambda value: re.sub(r"\[(?:TABLE|FIGURE)\d+\]", "", re.sub(r"\*\*", "", value))
    abstract = len(clean(abstract_match.group(1)).split())
    main = len(clean(main_match.group(1)).split())
    pre_references = len(clean(text.split("## Referencias")[0]).split())
    return {"abstract": abstract, "main_without_references": main, "all_before_references": pre_references}


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    counts = source_word_counts(text)
    if counts["abstract"] > 250:
        raise ValueError(f"Resumen excede 250 palabras: {counts['abstract']}")
    if counts["main_without_references"] > 3500:
        raise ValueError(f"Texto principal excede 3,500 palabras: {counts['main_without_references']}")
    if counts["all_before_references"] > 3500:
        raise ValueError("El texto con resumen y declaraciones excede 3,500 palabras.")

    document = Document()
    configure_document(document)
    document.core_properties.title = text.splitlines()[0].removeprefix("# ")
    mobile = QA_SCREEN
    if "[FIGURE4]" in text and not mobile.exists():
        raise FileNotFoundError("Falta captura real del prototipo: ejecutar scripts/ensanut-mobile-qa.js en modo oscuro.")

    lines = text.splitlines()
    seen_title = False
    in_references = False
    first_body_heading = False
    markers = []

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not seen_title:
            title = document.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(title, line[2:], size=11)
            seen_title = True
            continue
        if line.startswith("## ") and not first_body_heading:
            subtitle = document.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.paragraph_format.line_spacing = 1.5
            subtitle.paragraph_format.space_after = Pt(10)
            add_rich_text(subtitle, line[3:], size=11)
            first_body_heading = True
            continue
        if line.startswith("**Categoría propuesta:"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.15
            add_rich_text(paragraph, line, size=10)
            continue
        if line.startswith("**Tipo de documento:"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.15
            add_rich_text(paragraph, line, size=10)
            continue
        if line.startswith("**Estado:"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(12)
            run = paragraph.add_run("BORRADOR INTERNO. Pendiente de financiamiento y dictamen; no someter.")
            set_run_font(run, size=10, bold=True, color="9C2F3F")
            continue
        if line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_references = heading == "Referencias"
            document.add_paragraph(heading, style="Heading 1")
            continue
        if re.fullmatch(r"\[(TABLE|FIGURE)\d+\]", line):
            markers.append(line)
            continue
        if in_references and re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(4)
            add_rich_text(paragraph, line, size=11)
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        add_rich_text(paragraph, line, size=11)

    for marker in sorted(markers, key=lambda value: (not value.startswith("[TABLE"), value)):
        add_marker(document, marker, mobile)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
    for paragraph in paragraphs:
        paragraph.paragraph_format.line_spacing = 1.5
    document.save(OUTPUT)
    qa = {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "word_counts": counts,
        "references": len(re.findall(r"^\d+\. ", text, re.M)),
        "figures": sum(marker.startswith("[FIGURE") for marker in markers),
        "tables": sum(marker.startswith("[TABLE") for marker in markers),
        "main_source_sha256": __import__("hashlib").sha256(SOURCE.read_bytes()).hexdigest(),
        "blind_name_scan": {
            name: bool(re.search(name, text, re.I))
            for name in ("Guillermo", "Arianna", "Cristopher", "Gabriel", "Villafuerte", "Álvarez Sánchez")
        },
        "study_fingerprint": "cc911791bad871aca43e0ac677c4f44ccfa05343af94c07cf8faf5438c592b8f",
    }
    (HERE / "manuscript-fortalecido-qa.json").write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(qa, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
